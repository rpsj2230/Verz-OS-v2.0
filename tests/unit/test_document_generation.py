"""The tool that fills an approved template and keeps a Word file (M12.3.5), without a database.

The handler is driven through a row source that answers by what it is asked for and a store that
remembers what it was handed, so each property is asserted on the thing the handler decides: which
templates it will fill, what the file holds, who it is attributed to, and what it refuses. What a
statement admits is the document plane's own and is asserted in its own tests; the real rows are
the install check's.

Task ids: M12.3.5
"""

from __future__ import annotations

import asyncio
import io
from collections.abc import Awaitable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import Any

import docx
import pytest
from pydantic import ValidationError

from brain.console.agent_output import Artifact, ArtifactKind
from brain.core.entitlement import Capability, EntitlementSet, Grant
from brain.core.envelope import SideEffect, TypedResult
from brain.core.errors import Denied
from brain.core.scope import Scope
from brain.gate.catalogue import AgentCeiling, project
from brain.knowledge.document_tools import MAX_PASSAGES
from brain.knowledge.kinds import KnowledgeKind
from brain.knowledge.rows import RowQuery
from brain.knowledge.search import KNOWLEDGE_READ
from brain.ops.artifact_store import Produced
from brain.ops.retention import DataClass
from brain.tools.document_generation import (
    DOCUMENT_GENERATION_IS_NOT_AVAILABLE,
    GENERATE_CAPABILITY,
    GENERATE_DOCUMENT,
    GENERATED_DOCUMENT_READ,
    NO_APPROVED_TEMPLATE_IS_AVAILABLE,
    WORD_CONTENT_TYPE,
    BlockType,
    DocumentFormat,
    DocumentTools,
    GeneratedDocument,
    GenerateDocument,
    RunIdentity,
    TemplateFacts,
    agent_version_of,
    definition,
    facts_of,
    filled,
    handler,
    is_approved_template,
    layout,
    missing_values,
    placeholders_in,
    render,
)
from brain.tools.registry import ToolRegistry

#: Far from any wall clock, because nothing here is about the present. See CLAUDE.md on a fixture
#: with a date in it.
NOW = datetime(2999, 1, 1, 9, 0, tzinfo=UTC)
LONG_AGO = NOW - timedelta(days=400)
TEMPLATE = "doc_0123456789abcdef"
RUN = RunIdentity(agent_id="proposal_drafter", agent_version="sha256:abc", run_id="trace_1")
UPDATED = datetime(2019, 1, 1, 8, 0, tzinfo=UTC)

VERIFIED = {
    "kind": KnowledgeKind.TEMPLATE.value,
    "state": "published",
    "verified_by": "u_steward",
    "verified_at": LONG_AGO,
    "review_by": None,
}


def passage(ordinal: int, text: str, *, section: str = "") -> dict[str, Any]:
    return {
        "chunk_id": f"c_{ordinal}",
        "document_id": TEMPLATE,
        "ordinal": ordinal,
        "body": text,
        "title": "Proposal template",
        "section": section,
        "updated_at": UPDATED,
        "department": "web",
        "visibility": "department",
        "owner_id": "u_steward",
    }


BODY = (
    passage(0, "Proposal\nPrepared for {{client_name}}.", section="Proposal"),
    passage(1, "| Item | Price |\n| Build | {{price}} |"),
)


class Rows:
    """Answers the three questions the handler asks of the library, and remembers them."""

    def __init__(
        self,
        facts: Mapping[str, Any] | None = VERIFIED,
        passages: Sequence[Mapping[str, Any]] = BODY,
    ) -> None:
        self.facts = facts
        self.passages = tuple(passages)
        self.asked: list[RowQuery] = []

    async def rows(self, query: RowQuery) -> Sequence[Mapping[str, Any]]:
        self.asked.append(query)
        if query.entity == "department":
            return [{"slug": "web"}]
        if query.entity == "knowledge_item":
            return [] if self.facts is None else [dict(self.facts)]
        return list(self.passages)


@dataclass
class Shelf:
    """Where a file is kept: remembers what it was handed and mints the artifact it records."""

    kept: list[Produced] = field(default_factory=list)
    fail_with: Exception | None = None

    async def keep(self, produced: Produced, at: datetime) -> Artifact:
        if self.fail_with is not None:
            raise self.fail_with
        self.kept.append(produced)
        return Artifact(
            artifact_id=f"art_{len(self.kept)}",
            agent_id=produced.agent_id,
            kind=produced.kind,
            run_id=produced.run_id,
            agent_version=produced.agent_version,
            caller_id=produced.caller_id,
            entitlement_hash=produced.reach.ent_hash(),
            at=at,
            data_class=DataClass.PAYLOAD,
            bytes_stored=len(produced.body),
        )


def holding(*capabilities: str, who: str = "u_asker") -> EntitlementSet:
    return EntitlementSet(
        principal_id=who,
        grants=tuple(
            Grant(capability=Capability(value=value), scope=Scope.unrestricted())
            for value in capabilities
        ),
    )


ASKER = holding(KNOWLEDGE_READ.value, GENERATE_CAPABILITY.value)


def tools_over(rows: Rows, shelf: Shelf, *, on: bool = True) -> DocumentTools:
    async def is_on() -> bool:
        return on

    return DocumentTools(records=rows, keep=shelf.keep, is_on=is_on)


def settle[T](pending: Awaitable[T]) -> T:
    async def wait() -> T:
        return await pending

    return asyncio.run(wait())


def generate(
    rows: Rows,
    shelf: Shelf,
    request: GenerateDocument,
    *,
    entitlement: EntitlementSet = ASKER,
    produced_for: RunIdentity | None = RUN,
    on: bool = True,
) -> TypedResult[GeneratedDocument]:
    call = handler(tools_over(rows, shelf, on=on))
    return settle(call(request, entitlement=entitlement, now=NOW, produced_for=produced_for))


def asking(**values: str) -> GenerateDocument:
    return GenerateDocument(template_id=TEMPLATE, title="Proposal for Acme", values=values)


GOOD = {"client_name": "Acme Ltd", "price": "1,200"}


def read_back(body: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(body))


def words_of(body: bytes) -> str:
    document = read_back(body)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([one.text for one in document.paragraphs] + cells)


# ------------------------------------------------------------------ the positive case
def test_an_approved_template_is_filled_and_kept_for_the_person_who_asked() -> None:
    """**The positive case every refusal below needs beside it.** A verified, published template
    is read, its placeholders take the values supplied (in a paragraph and in a table cell), and
    the Word file is kept as a document artifact attributed to the person asking and to the run,
    with the template named as what it drew on and the result naming the artifact and no content.

    Delete this and a handler that refuses everything passes every refusal in this file."""
    rows, shelf = Rows(), Shelf()

    result = generate(rows, shelf, asking(**GOOD))

    (kept,) = shelf.kept
    text = words_of(kept.body)
    assert "Prepared for Acme Ltd." in text and "1,200" in text and "{{" not in text
    assert (kept.kind, kept.content_type) == (ArtifactKind.DOCUMENT, WORD_CONTENT_TYPE)
    assert (kept.agent_id, kept.run_id, kept.agent_version) == (
        RUN.agent_id,
        RUN.run_id,
        RUN.agent_version,
    )
    assert kept.caller_id == ASKER.principal_id
    assert kept.provenance.knowledge_items == (TEMPLATE,)
    (record,) = result.records
    assert (record.artifact_id, record.title, record.format) == (
        "art_1",
        "Proposal for Acme",
        "word",
    )
    assert record.size_bytes == len(kept.body)
    assert set(GeneratedDocument.model_fields) == {
        "entity",
        "id",
        "artifact_id",
        "title",
        "format",
        "size_bytes",
    }


def test_the_file_is_kept_at_the_run_reach_and_remembers_the_grant_it_drew_on() -> None:
    """The reach recorded is the one the handler was handed, which the gate already narrowed by
    the agent's ceiling, and the knowledge grant it read the template under is written onto the
    record so a later download is checked against it. Delete this and a file can be kept at some
    other reach, or with nothing to say what it was drawn from, and a person whose reach has since
    narrowed could still fetch what their old reach built."""
    shelf = Shelf()

    generate(Rows(), shelf, asking(**GOOD))

    (kept,) = shelf.kept
    assert kept.reach == ASKER
    assert [one.capability for one in kept.drew_on] == [KNOWLEDGE_READ]


def test_the_file_is_labelled_with_the_shortest_lived_thing_it_was_built_from() -> None:
    """The template is knowledge and the agent's words are a run's payload, and the artifact is
    kept under the narrower of the two. Delete this and a document built from a thirty-day payload
    is labelled as if it outlived it, which the artifact module's rule on a copy that outlives its
    source forbids."""
    from brain.console.agent_output import retention_class_for

    shelf = Shelf()

    generate(Rows(), shelf, asking(**GOOD))

    assert retention_class_for(shelf.kept[0].inputs) is DataClass.PAYLOAD
    assert {one.data_class for one in shelf.kept[0].inputs} == {
        DataClass.LEARNED,
        DataClass.PAYLOAD,
    }


# ------------------------------------------------------------------ what is a template
@pytest.mark.parametrize(
    "facts",
    [
        {**VERIFIED, "verified_by": "", "verified_at": None},
        {**VERIFIED, "state": "draft"},
        {**VERIFIED, "state": "superseded"},
        {**VERIFIED, "state": "archived"},
        {**VERIFIED, "kind": KnowledgeKind.SOP.value},
        {**VERIFIED, "kind": None},
        None,
    ],
    ids=["unverified", "draft", "superseded", "archived", "other-kind", "no-kind", "unreadable"],
)
def test_a_template_nobody_vouched_for_is_never_filled_and_every_way_is_one_sentence(
    facts: Mapping[str, Any] | None,
) -> None:
    """Seven ways a reference is not an approved template, which is a document that is not one
    and a document the person may not read, each refused with the same sentence and nothing kept.
    Delete this and a draft or an unverified document can be filled, or the refusal can say which
    way it failed and so say which documents exist."""
    shelf = Shelf()

    with pytest.raises(Denied) as refused:
        generate(Rows(facts=facts), shelf, asking(**GOOD))

    assert str(refused.value) == NO_APPROVED_TEMPLATE_IS_AVAILABLE
    assert shelf.kept == []


def test_a_template_past_its_review_date_is_still_a_verified_one() -> None:
    """Verified and overdue for review is the badge's `due`, which still says somebody vouched.
    Delete this and the day a review lapses silently stops every document an agent makes from that
    template, with nothing to say why."""
    shelf = Shelf()
    overdue = {**VERIFIED, "review_by": NOW - timedelta(days=1)}

    generate(Rows(facts=overdue), shelf, asking(**GOOD))

    assert len(shelf.kept) == 1


def test_approval_reads_the_same_badge_an_answer_is_shown_with() -> None:
    """The approval is `brain.knowledge.item.badge`'s, not a second reading of verification.
    Delete this and the two can disagree, and a document is filled from a template an answer's
    own badge calls unverified."""
    plain = facts_of(VERIFIED)
    assert plain is not None and is_approved_template(plain, now=NOW)
    half = TemplateFacts(
        kind="template",
        state=plain.state,
        verified_by="",
        verified_at=None,
        review_by=None,
    )
    assert not is_approved_template(half, now=NOW)


def test_a_row_whose_state_is_not_one_the_library_has_is_not_a_template() -> None:
    """A state outside the closed four reads as no facts rather than raising. Delete this and a
    row somebody wrote by hand turns a refusal into a failed run."""
    assert facts_of({**VERIFIED, "state": "made_up"}) is None


def test_a_template_the_read_cut_short_is_refused_rather_than_filled_in_part() -> None:
    """A template as long as the read's own limit may have been cut, and a file made from half a
    template is a wrong document with nothing to say so. Delete this and the first fifty passages
    of a longer template are quietly the whole of it."""
    cut = [passage(i, f"Line {i}") for i in range(MAX_PASSAGES)]
    shelf = Shelf()

    with pytest.raises(Denied) as refused:
        generate(Rows(passages=cut), shelf, asking())

    assert str(refused.value) == NO_APPROVED_TEMPLATE_IS_AVAILABLE
    assert shelf.kept == []


def test_a_template_with_nothing_readable_in_it_is_refused() -> None:
    """Facts that say approved and no passage the caller may read. Delete this and a template
    whose chunks are not visible produces a file with only a title."""
    with pytest.raises(Denied) as refused:
        generate(Rows(passages=()), Shelf(), asking())

    assert str(refused.value) == NO_APPROVED_TEMPLATE_IS_AVAILABLE


# ------------------------------------------------------------------ the values
def test_a_placeholder_with_no_value_is_refused_by_name_and_nothing_is_kept() -> None:
    """The one refusal that names anything: the template's own placeholders, which the person
    asking may read. Delete this and a file is kept with a literal `{{price}}` in it."""
    shelf = Shelf()

    with pytest.raises(Denied) as refused:
        generate(Rows(), shelf, asking(client_name="Acme Ltd"))

    assert "price" in str(refused.value) and "client_name" not in str(refused.value)
    assert shelf.kept == []


def test_a_value_that_names_another_placeholder_is_written_as_text() -> None:
    """One pass, so a value holding `{{price}}` is that text in the document and does not pull
    the price in. Delete this and one value can reach into another's place.
    See `A_VALUE_IS_TEXT_AND_NEVER_A_TEMPLATE`."""
    shelf = Shelf()

    generate(Rows(), shelf, asking(client_name="{{price}}", price="9"))

    text = words_of(shelf.kept[0].body)
    assert "Prepared for {{price}}." in text and "| 9 |" not in text


def test_a_value_the_template_does_not_name_is_ignored() -> None:
    """Extra values are not written anywhere. Delete this and a model's stray key can surface in a
    document the template never gave it a place in."""
    shelf = Shelf()

    generate(Rows(), shelf, asking(**GOOD, secret_note="not for the page"))

    assert "not for the page" not in words_of(shelf.kept[0].body)


def test_characters_a_word_file_cannot_hold_are_dropped_not_fatal() -> None:
    """A model can write a control character, and the library refuses one. Delete this and one
    stray byte turns a good request into a failed run."""
    shelf = Shelf()

    generate(Rows(), shelf, asking(client_name="Ac\x00me\x0b Ltd", price="1"))

    assert "Prepared for Acme Ltd." in words_of(shelf.kept[0].body)


@pytest.mark.parametrize(
    "bad",
    [
        {"template_id": "../etc", "title": "t"},
        {"template_id": TEMPLATE, "title": ""},
        {"template_id": TEMPLATE, "title": "t", "format": "powerpoint"},
        {"template_id": TEMPLATE, "title": "t", "format": "excel"},
        {"template_id": TEMPLATE, "title": "t", "values": {"Client Name": "x"}},
        {"template_id": TEMPLATE, "title": "t", "values": {"a": "x" * 20_001}},
        {"template_id": TEMPLATE, "title": "t", "values": {f"v{i}": "x" for i in range(101)}},
        {
            "template_id": TEMPLATE,
            "title": "t",
            "values": {f"v{i}": "x" * 20_000 for i in range(11)},
        },
        {"template_id": TEMPLATE, "title": "t", "something_else": "x"},
    ],
    ids=[
        "reference",
        "title",
        "deck",
        "spreadsheet",
        "name",
        "long-value",
        "many-values",
        "long-total",
        "unknown-field",
    ],
)
def test_a_request_the_schema_refuses_never_reaches_the_handler(bad: dict[str, Any]) -> None:
    """Nine malformed requests, and a slide deck and a spreadsheet among them, each refused when
    the request is built. Delete this and the closed format list is a comment, and a model can
    ask for what the product does not hold."""
    with pytest.raises(ValidationError):
        GenerateDocument.model_validate(bad)
    assert [one.value for one in DocumentFormat] == ["word"]


# ------------------------------------------------------------------ the switch and the run
def test_a_tool_switched_off_produces_nothing_and_says_it_is_unavailable() -> None:
    """The switch ships off and is asked at the call. Delete this and a registry built before an
    administrator switched it on keeps producing files after they switched it off."""
    rows, shelf = Rows(), Shelf()

    with pytest.raises(Denied) as refused:
        generate(rows, shelf, asking(**GOOD), on=False)

    assert str(refused.value) == DOCUMENT_GENERATION_IS_NOT_AVAILABLE
    assert shelf.kept == [] and rows.asked == []


def test_a_call_with_no_agent_behind_it_produces_nothing() -> None:
    """An automation's tool calls bind no agent today, and a file attributed to nobody is not
    made. Delete this and a scheduled run keeps files with an agent id of nothing."""
    shelf = Shelf()

    with pytest.raises(Denied) as refused:
        generate(Rows(), shelf, asking(**GOOD), produced_for=None)

    assert str(refused.value) == DOCUMENT_GENERATION_IS_NOT_AVAILABLE
    assert shelf.kept == []


def test_a_store_that_cannot_keep_a_file_is_an_unavailable_tool_and_not_a_failed_run() -> None:
    """The artifact store refusing, or having no bucket, answers as an unavailable tool does.
    Delete this and the run ends on an exception the person is shown."""
    from brain.ops.storage import StorageError

    shelf = Shelf(fail_with=StorageError("no store"))

    with pytest.raises(Denied) as refused:
        generate(Rows(), shelf, asking(**GOOD))

    assert str(refused.value) == DOCUMENT_GENERATION_IS_NOT_AVAILABLE


def test_nothing_is_read_for_a_caller_who_reaches_no_knowledge() -> None:
    """With no grant on the document plane the reach is none and the template is not asked for.
    Delete this and a caller with no knowledge grant is sent through the template reads anyway."""
    rows, shelf = Rows(), Shelf()

    with pytest.raises(Denied) as refused:
        generate(rows, shelf, asking(**GOOD), entitlement=holding(GENERATE_CAPABILITY.value))

    assert str(refused.value) == NO_APPROVED_TEMPLATE_IS_AVAILABLE
    assert all(one.entity == "department" for one in rows.asked)
    assert shelf.kept == []


# ------------------------------------------------------------------ the layout and the file
def test_a_template_is_laid_out_as_headings_paragraphs_and_tables() -> None:
    """A line the section names that opens a passage is a heading, a Markdown heading is one, a
    run of pipe rows is a table and the rest are paragraphs. Delete this and a template's
    structure collapses into one run of text."""
    from brain.knowledge.document_tools import KnowledgePassage

    def one(text: str, section: str = "") -> KnowledgePassage:
        return KnowledgePassage(
            entity="knowledge", id="c", document_id=TEMPLATE, document=text, section=section
        )

    blocks = layout(
        [
            one("Scope\nWe build it.\n\nWe test it.", section="Scope"),
            one("## Terms\nPay in thirty days."),
            one("| A | B |\n| 1 | 2 |"),
        ]
    )

    assert [(b.kind, b.lines or b.rows) for b in blocks] == [
        (BlockType.HEADING, ("Scope",)),
        (BlockType.PARAGRAPH, ("We build it.",)),
        (BlockType.PARAGRAPH, ("We test it.",)),
        (BlockType.HEADING, ("Terms",)),
        (BlockType.PARAGRAPH, ("Pay in thirty days.",)),
        (BlockType.TABLE, (("A", "B"), ("1", "2"))),
    ]
    assert blocks[3].level == 2


def test_a_line_that_only_repeats_the_section_later_in_a_passage_is_not_a_heading() -> None:
    """Only the line that opens a passage is read as its heading. Delete this and a sentence that
    happens to equal the section name becomes a heading in the middle of a paragraph."""
    from brain.knowledge.document_tools import KnowledgePassage

    blocks = layout(
        [
            KnowledgePassage(
                entity="knowledge",
                id="c",
                document_id=TEMPLATE,
                document="We build it.\nScope",
                section="Scope",
            )
        ]
    )

    assert [b.kind for b in blocks] == [BlockType.PARAGRAPH]


def test_placeholders_are_found_once_each_in_the_order_they_first_appear() -> None:
    """Delete this and a template that names one placeholder five times asks for it five times,
    or asks for them in an order a person cannot follow."""
    from brain.tools.document_generation import Block

    blocks = (
        Block(kind=BlockType.PARAGRAPH, lines=("{{b}} and {{a}}", "{{ b }}")),
        Block(kind=BlockType.TABLE, rows=(("{{c}}", "x"),)),
    )

    assert placeholders_in(blocks) == ("b", "a", "c")
    assert missing_values(blocks, {"a": "1"}) == ("b", "c")
    assert filled("{{ b }}!", {"b": "yes"}) == "yes!"
    assert filled("{{Upper}}", {"upper": "no"}) == "{{Upper}}"


def test_the_file_carries_no_name_of_an_author_and_the_runs_own_date() -> None:
    """The library's template carries an author and a date of its own. Delete this and every file
    an agent makes names somebody as its author."""
    body = render("Proposal", (), {}, now=NOW)

    properties = read_back(body).core_properties
    assert properties.author == "" and properties.last_modified_by == ""
    assert properties.title == "Proposal"
    assert properties.created == NOW.replace(tzinfo=None) or properties.created == NOW


def test_the_agent_version_follows_what_the_agent_is_and_not_whether_it_is_enabled() -> None:
    """Disabling an agent changes neither what it is nor what it made. Delete this and a version
    that moves with a toggle records a different agent for the same file."""
    from brain.agents.model import AgentAudience, AgentAuthority, AgentRecord
    from brain.knowledge.visibility import Visibility

    record = AgentRecord(
        agent_id="proposal_drafter",
        display_name="Proposal drafter",
        persona="Drafts proposals.",
        audience=AgentAudience(owner_id="u_o", level=Visibility.COMPANY),
        authority=AgentAuthority(
            allowed_tools=frozenset(), required_tools=frozenset(), max_side_effect=SideEffect.NONE
        ),
        created_by="u_o",
    )
    paused = record.model_copy(update={"disabled_at": NOW})
    changed = record.model_copy(update={"persona": "Drafts something else."})

    assert agent_version_of(record) == agent_version_of(paused)
    assert agent_version_of(record) != agent_version_of(changed)
    assert agent_version_of(record).startswith("sha256:")


# ------------------------------------------------------------------ the registration
def test_the_tool_is_offered_only_to_an_agent_whose_ceiling_names_it() -> None:
    """A caller who holds the capability still does not see the tool through an agent whose
    ceiling does not name it, and an agent that names it shows nothing to a caller without the
    capability. Delete this and the tool can reach a run whose agent was never given it.

    The catalogue is `E(caller)` against the ceiling, projected by the one projector, so this asks
    nothing of a second implementation of the rule."""
    registry = ToolRegistry()
    from brain.tools.document_generation import register_document_generation

    register_document_generation(registry, tools_over(Rows(), Shelf()))
    other = definition().model_copy(update={"name": "documents.read_template", "description": "x"})
    registry.register(other, _echo)

    named = AgentCeiling(agent_id="a", allowed_tools=frozenset({GENERATE_DOCUMENT}))
    unnamed = AgentCeiling(agent_id="a", allowed_tools=frozenset({"documents.read_template"}))

    assert project(registry, ASKER, named, now=NOW).names == (GENERATE_DOCUMENT,)
    assert GENERATE_DOCUMENT not in project(registry, ASKER, unnamed, now=NOW).names
    assert project(registry, holding(KNOWLEDGE_READ.value), named, now=NOW).names == ()


async def _echo(
    request: GenerateDocument, *, entitlement: EntitlementSet
) -> TypedResult[GeneratedDocument]:
    del request, entitlement
    return TypedResult[GeneratedDocument]()


def test_the_definition_is_a_read_that_asks_for_its_own_capability() -> None:
    """Declared as a read under `invoke:generated_document`, delegated, with the format closed in
    its schema. Delete this and the tool can be declared as an effect nothing offers to a run,
    or under a capability nobody holds."""
    declared = definition()

    assert declared.name == GENERATE_DOCUMENT
    assert declared.required_capability == GENERATE_CAPABILITY.value
    assert declared.side_effect is SideEffect.NONE
    assert declared.args_schema["$defs"]["DocumentFormat"]["enum"] == ["word"]
    assert declared.entity == "generated_document"


# ------------------------------------------------------------------ the words that must not move
def test_the_words_a_run_and_an_administrator_read_are_the_ones_written_here() -> None:
    """The refusal sentence, the capability, the tool's name and the file's media type are each
    asserted against the text, because a test that compares a constant with itself passes for every
    value it could hold. Delete this and the sentence can start to say which of five things was
    wrong, or the capability can move to one nobody grants, with every other test green."""
    assert NO_APPROVED_TEMPLATE_IS_AVAILABLE == (
        "no approved template by that reference is available to this agent"
    )
    assert (
        DOCUMENT_GENERATION_IS_NOT_AVAILABLE == "document generation is not available to this run"
    )
    assert GENERATE_DOCUMENT == "documents.generate_document"
    assert GENERATE_CAPABILITY.value == "invoke:generated_document"
    assert GENERATED_DOCUMENT_READ.value == "read:generated_document"
    assert WORD_CONTENT_TYPE == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_what_a_run_is_told_of_the_file_is_read_under_the_capability_for_reading_it() -> None:
    """The result is redacted with the tool's own policy: a run that may read the result sees the
    reference, title, format and size, and one that was handed only the permission to call the
    tool sees nothing of it. Delete this and a result of an entity nobody classified is withheld
    whole, which is what a model that has just made a file would report as having made nothing."""
    from brain.core.redaction import redact
    from brain.tools.own_results import OWN_RESULT_POLICIES

    result = TypedResult[GeneratedDocument](
        records=(
            GeneratedDocument(
                entity="generated_document",
                id="art_1",
                artifact_id="art_1",
                title="Proposal",
                format="word",
                size_bytes=10,
            ),
        ),
        source="documents",
    )
    policy = OWN_RESULT_POLICIES["generated_document"]

    seen = redact(
        result, entitlement=holding(GENERATED_DOCUMENT_READ.value), policy=policy, now=NOW
    )
    blind = redact(result, entitlement=holding(GENERATE_CAPABILITY.value), policy=policy, now=NOW)

    assert next(dict(one) for one in seen.payload.records)["artifact_id"] == "art_1"
    assert [dict(one) for one in blind.payload.records] == []


def test_a_registry_offers_the_tool_only_where_it_was_bound_to_a_template_source_and_a_store() -> (
    None
):
    """`build_registry` registers it when handed the binding and not otherwise. Delete this and a
    tool bound to nothing can sit in the catalogue, offered to a model and unable to answer."""
    from brain.tools.startup import build_registry

    bound = build_registry(
        source="demo",
        records=Rows(),
        documents=tools_over(Rows(), Shelf()),
    )
    unbound = build_registry(source="demo", records=Rows())

    assert bound.has(GENERATE_DOCUMENT)
    assert not unbound.has(GENERATE_DOCUMENT)


def test_a_run_binds_the_agent_it_is_and_an_automation_binds_nobody() -> None:
    """A run's tool caller hands the tool the agent and run a file is attributed to, and a caller
    with nothing bound hands it none, so the tool refuses. Delete this and a run can reach the tool
    with its agent left out, or an automation can produce files attributed to nobody."""
    from brain.api_routes import RunToolCaller
    from brain.automation_routes import RegistryToolCaller
    from brain.tools.document_generation import register_document_generation

    registry = ToolRegistry()
    shelf = Shelf()
    register_document_generation(registry, tools_over(Rows(), shelf))
    arguments = {"template_id": TEMPLATE, "title": "t", "values": GOOD}

    running = RunToolCaller(registry, agent_id=RUN.agent_id, produced_for=RUN)
    settle(running.call(tool=definition(), arguments=arguments, entitlement=ASKER, now=NOW))  # type: ignore[arg-type]
    assert [one.agent_id for one in shelf.kept] == [RUN.agent_id]

    unattended = RegistryToolCaller(registry)
    with pytest.raises(Denied):
        settle(unattended.call(tool=definition(), arguments=arguments, entitlement=ASKER, now=NOW))  # type: ignore[arg-type]
    assert len(shelf.kept) == 1
