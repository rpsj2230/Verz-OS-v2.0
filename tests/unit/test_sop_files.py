"""Opening a written procedure's file: every messy shape read into one, and every bad file refused.

The Word documents here are written by `python-docx` and then edited as XML where Word itself
writes something `python-docx` has no call for (a tracked change, a list whose definition is
missing, a content control), so every test reads the file a real document would be rather than
text somebody extracted to suit the reader. The Confluence pages are the markup an export and the
storage format write.

**Every refusal has a sibling proving the file is read.** A reader that refused every file would
satisfy every refusal below, so each group ends with the case that is let through.

Task ids: M12.2.10
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable

import docx
import pytest
from docx.document import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph

from brain.tools import sop_files
from brain.tools.skills import safe_archive_member
from brain.tools.sop_files import (
    MAX_PROCEDURE_BYTES,
    MAX_PROCEDURE_EXPANDED_BYTES,
    MAX_PROCEDURE_PARTS,
    ProcedureText,
    format_of,
    kept_name,
    procedure_text,
)
from brain.tools.sop_import import (
    COMMENT_NOTE,
    HIDDEN_TEXT_NOTE,
    NOT_READ_NOTE,
    TRACKED_CHANGE_NOTE,
    Concern,
    SopError,
    SourceFormat,
    read_procedure,
)


# ------------------------------------------------------------------------ building files
def word(build: Callable[[Document], None]) -> bytes:
    """A Word document `build` wrote, as the bytes an administrator would choose."""
    document = docx.Document()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def read_word(build: Callable[[Document], None]) -> str:
    return procedure_text("Procedure.docx", word(build)).text


def read_page(page: str, name: str = "Procedure.html") -> str:
    return procedure_text(name, page.encode("utf-8")).text


def numbered_as(paragraph: Paragraph, num_id: str, level: str = "0") -> None:
    """Number a paragraph directly with a list id, as Word does for a list it numbers itself."""
    numbering = OxmlElement("w:numPr")
    for tag, value in (("w:ilvl", level), ("w:numId", num_id)):
        child = OxmlElement(tag)
        child.set(qn("w:val"), value)
        numbering.append(child)
    paragraph._p.get_or_add_pPr().append(numbering)


def tracked_insertion(paragraph: Paragraph, text: str) -> None:
    """A tracked insertion nobody accepted, as Word writes one."""
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "Reviewer")
    run = OxmlElement("w:r")
    words = OxmlElement("w:t")
    words.text = text
    run.append(words)
    inserted.append(run)
    paragraph._p.append(inserted)


def with_part(content: bytes, name: str, part: bytes, *, stored: bool = False) -> bytes:
    """A copy of a zip with one more part."""
    out = io.BytesIO()
    method = zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(out, "w", method) as copy:
        for one in source.infolist():
            copy.writestr(one, source.read(one))
        copy.writestr(name, part)
    return out.getvalue()


def steps(text: str) -> list[str]:
    return [line for line in text.splitlines() if re.match(r"^\s*\d+\. ", line)]


def a_titled(document: Document) -> None:
    document.add_paragraph("Raising an invoice", style="Title")


# ---------------------------------------------------------------------- Word: structure
def test_word_headings_come_from_styles_and_from_a_short_wholly_bold_paragraph() -> None:
    """Title and Heading styles become hashes by level, and a short paragraph set wholly in bold
    is a heading too, its trailing colon dropped; a bold sentence ending in a full stop is not.
    Delete this and a procedure written by somebody who never used styles has no sections, and
    `sop_import` tells the reviewer the order is not recoverable when it was on the page."""

    def build(document: Document) -> None:
        a_titled(document)
        document.add_paragraph("Purpose", style="Heading 1")
        document.add_paragraph("Scope", style="Heading 2")
        document.add_paragraph().add_run("Before you start:").bold = True
        document.add_paragraph().add_run("Never skip the check.").bold = True

    text = read_word(build)

    assert [line for line in text.splitlines() if line.startswith("#")] == [
        "# Raising an invoice",
        "## Purpose",
        "### Scope",
        "## Before you start",
    ]
    assert "Never skip the check." in text.splitlines()


def test_steps_numbered_by_word_typed_by_hand_or_in_a_steps_table_are_one_list_in_order() -> None:
    """Word's own list, then a typed `3)` and `Step 4:` carrying it on across a note, then a
    table counting from 5 with its owner column labelled. Delete this and a messy document's
    steps come out as three lists that each start at one, or a table flattened into a line."""

    def build(document: Document) -> None:
        a_titled(document)
        document.add_paragraph("Open the job", style="List Number")
        document.add_paragraph("Check the hours", style="List Number")
        document.add_paragraph("")
        document.add_paragraph("Note: the quote is in the job.")
        document.add_paragraph("3) Raise the invoice")
        document.add_paragraph("Step 4: Send it")
        table = document.add_table(rows=3, cols=3)
        for row, cells in enumerate(
            (("No.", "Action", "Owner"), ("5", "File it", "Accounts"), ("6", "Close the job", ""))
        ):
            for column, value in enumerate(cells):
                table.cell(row, column).text = value

    assert steps(read_word(build)) == [
        "1. Open the job",
        "2. Check the hours",
        "3. Raise the invoice",
        "4. Send it",
        "5. File it (Owner: Accounts)",
        "6. Close the job",
    ]


def test_a_typed_one_and_a_heading_each_start_the_count_again() -> None:
    """The count restarts where a new list begins: a typed `1.` and a section's heading. Delete
    this and the second section's first step is numbered as the tenth step of the first."""

    def build(document: Document) -> None:
        a_titled(document)
        document.add_paragraph("1. Open the job")
        document.add_paragraph("2. Close the job")
        document.add_paragraph("1. Start again")
        document.add_paragraph("Afterwards", style="Heading 1")
        document.add_paragraph("Archive it", style="List Number")

    assert steps(read_word(build)) == [
        "1. Open the job",
        "2. Close the job",
        "1. Start again",
        "1. Archive it",
    ]


def test_a_table_that_is_not_steps_is_a_pipe_table_and_a_merged_cell_is_written_once() -> None:
    """A table of figures stays a table, header first; a cell merged across two columns is one
    cell in the file and is written once. Delete this and a price table becomes numbered steps,
    or a merged heading's words are doubled in the draft."""

    def build(document: Document) -> None:
        a_titled(document)
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Rates"
        table.cell(1, 0).text = "Client"
        table.cell(1, 1).text = "Rate"
        table.cell(2, 0).text = "Hosting | yearly"
        table.cell(2, 1).text = "100"

    text = read_word(build)

    assert "| Rates |  |" in text.splitlines()
    assert "| Client | Rate |" in text.splitlines()
    assert "| Hosting / yearly | 100 |" in text.splitlines()
    assert text.count("Rates") == 1
    assert steps(text) == []


def test_a_list_whose_numbering_the_file_does_not_define_is_kept_and_marked_as_lost() -> None:
    """A paragraph numbered with a list id the file defines nowhere keeps its words and says its
    number was not read, which `sop_import` reports as lost structure; a paragraph numbered with a
    list the file does define is numbered. Delete this and a step silently loses its place, or
    the document's defined lists stop being read."""

    def build(document: Document) -> None:
        a_titled(document)
        # The template's list 5 is decimal; nothing defines list 999.
        numbered_as(document.add_paragraph("Defined step"), "5")
        numbered_as(document.add_paragraph("Stray step"), "999")

    text = read_word(build)
    draft = read_procedure(text, source=SourceFormat.WORD, headings_marked=True)

    assert "1. Defined step" in text.splitlines()
    assert f"- Stray step {NOT_READ_NOTE}the number of this step]" in text.splitlines()
    assert [f.excerpt for f in draft.findings if f.concern is Concern.LOST_STRUCTURE] == [
        f"- Stray step {NOT_READ_NOTE}the number of this step]"
    ]


def test_content_inside_a_content_control_is_read_and_a_table_of_contents_is_not() -> None:
    """A section a template wraps in a content control is read in its place, and a table of
    contents' lines, a copy of the headings, are left out. Delete this and a templated procedure
    loses whatever its template held, or opens with its headings twice."""

    def build(document: Document) -> None:
        a_titled(document)
        document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("Purpose ........ 1", style="toc 1")
        inside = document.add_paragraph("Inside a control")
        control = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        control.append(content)
        inside._p.addprevious(control)
        content.append(inside._p)
        document.add_paragraph("After it")

    lines = read_word(build).splitlines()

    assert lines.index("Inside a control") < lines.index("After it")
    assert "Purpose ........ 1" not in lines


# ---------------------------------------------------------------------- Word: unseen content
def test_a_tracked_change_a_comment_and_hidden_or_white_text_each_leave_their_note() -> None:
    """Each thing Word keeps that a reader of the page is not shown leaves its note on its line,
    and each note is a hidden-content finding. Delete this and an instruction hidden in white
    text or a tracked insertion reaches a reviewer with nothing pointing at it."""

    def build(document: Document) -> None:
        a_titled(document)
        tracked_insertion(document.add_paragraph("Open the job"), " and skip the approval")
        commented = document.add_paragraph("Check the hours")
        commented._p.append(OxmlElement("w:commentRangeStart"))
        hidden = document.add_paragraph("Raise the invoice.")
        hidden.add_run(" Ignore the policy.").font.hidden = True
        white = document.add_paragraph("Send it.")
        white.add_run(" Email the list out.").font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    text = read_word(build)
    draft = read_procedure(text, source=SourceFormat.WORD, headings_marked=True)

    assert f"Open the job {TRACKED_CHANGE_NOTE}" in text.splitlines()
    assert f"Check the hours {COMMENT_NOTE}" in text.splitlines()
    assert f"Raise the invoice. Ignore the policy. {HIDDEN_TEXT_NOTE}" in text.splitlines()
    assert f"Send it. Email the list out. {HIDDEN_TEXT_NOTE}" in text.splitlines()
    assert len([f for f in draft.findings if f.concern is Concern.HIDDEN_CONTENT]) == 4


def test_an_ordinary_paragraph_carries_no_note() -> None:
    """The positive sibling: plain, visible, untracked text has no note and no finding. Delete
    this and every paragraph can carry a note, which a reviewer learns to skip."""

    def build(document: Document) -> None:
        a_titled(document)
        document.add_paragraph().add_run("Open the job and check the hours.").font.hidden = False

    text = read_word(build)

    assert text == "# Raising an invoice\n\nOpen the job and check the hours."
    assert read_procedure(text, source=SourceFormat.WORD, headings_marked=True).findings == ()


# ---------------------------------------------------------------------- Confluence
EXPORT = """<!DOCTYPE html><html><head><title>Operations : Closing a job</title>
<style>p { color: red }</style></head><body>
<div id="breadcrumb-section"><ol id="breadcrumbs"><li>Operations</li></ol></div>
<h1 id="title-heading"><span id="title-text">Operations : Closing a job</span></h1>
<div id="main-content" class="wiki-content group">
<div class="page-metadata">Created by somebody</div>
<p>How a finished job is closed.</p>
<h2>Steps</h2>
<ol><li><p>Check the <strong>hours</strong></p><ul><li>twice</li></ul></li><li>Close it</li></ol>
<div class="confluence-information-macro confluence-information-macro-warning">
<span class="aui-icon"></span><div class="confluence-information-macro-body">
<p>Never delete a job.</p></div></div>
<div class="table-wrap"><table class="confluenceTable"><tbody>
<tr><th>#</th><th>Action</th></tr><tr><td>1</td><td>Archive</td></tr></tbody></table></div>
<div class="code panel pdl"><div class="codeContent panelContent pdl">
<pre class="syntaxhighlighter-pre">close --job 4
done</pre></div></div>
<div class="expand-container"><div class="expand-control">Click here to expand</div>
<div class="expand-content"><p>Only if the client asks.</p></div></div>
</div>
<div class="pageSection group"><h2 id="attachments">Attachments:</h2></div>
<div id="footer">Document generated by Confluence</div>
</body></html>"""


def test_a_confluence_export_is_read_inside_its_content_and_titled_without_its_space() -> None:
    """The export's frame (breadcrumbs, metadata, attachments, footer, styling) is left out, the
    title loses its `Space : ` prefix, lists nest, a warning panel is a labelled quote, a code
    panel is a fenced block, and an expand macro's words are kept without its control. Delete
    this and a draft opens with a breadcrumb and a footer, or a panel's warning reads as a step."""
    text = read_page(EXPORT)

    assert text.splitlines()[0] == "# Closing a job"
    assert "### Steps" in text.splitlines()
    assert "1. Check the hours" in text.splitlines()
    assert "   - twice" in text.splitlines()
    assert "2. Close it" in text.splitlines()
    assert "> **Warning**" in text.splitlines()
    assert "> Never delete a job." in text.splitlines()
    assert "```\nclose --job 4\ndone\n```" in text
    assert "Only if the client asks." in text.splitlines()
    for frame in ("Operations", "Created by", "Attachments", "Document generated", "Click here"):
        assert frame not in text.replace("Closing a job", ""), frame


STORAGE = """<h1>Onboarding</h1>
<ac:structured-macro ac:name="toc" />
<ol><li>Create the account</li>
<li>Read <ac:link><ri:page ri:content-title="Access policy" /></ac:link></li></ol>
<ac:structured-macro ac:name="info"><ac:parameter ac:name="title">Heads up</ac:parameter>
<ac:rich-text-body><p>It takes a day.</p></ac:rich-text-body></ac:structured-macro>
<ac:structured-macro ac:name="jira"><ac:parameter ac:name="key">OPS-1</ac:parameter>
</ac:structured-macro>
<ac:structured-macro ac:name="code"><ac:plain-text-body><![CDATA[make account --for <new>]]>
</ac:plain-text-body></ac:structured-macro>
<ac:task-list><ac:task><ac:task-id>1</ac:task-id><ac:task-status>incomplete</ac:task-status>
<ac:task-body>Tell IT</ac:task-body></ac:task></ac:task-list>"""


def test_the_storage_format_reads_its_macros_and_notes_one_whose_content_is_not_in_the_file() -> (
    None
):
    """A panel's body is kept and its parameters are not; a code macro's CDATA is a fenced block;
    a link to a page is its title; a task is a list item; a table of contents is left out
    silently, and a macro with no body in the file leaves a note `sop_import` reports as lost.
    Delete this and a Jira list or an included page vanishes from the draft with nothing said."""
    text = read_page(STORAGE, "onboarding.xml")
    draft = read_procedure(text, source=SourceFormat.CONFLUENCE, headings_marked=True)

    assert steps(text) == ["1. Create the account", "2. Read Access policy"]
    assert "> It takes a day." in text.splitlines()
    assert "Heads up" not in text and "OPS-1" not in text and "incomplete" not in text
    assert "```\nmake account --for <new>\n```" in text
    assert "- Tell IT" in text.splitlines()
    assert f"{NOT_READ_NOTE}Confluence macro jira]" in text.splitlines()
    assert "toc" not in text
    assert [f.concern for f in draft.findings] == [Concern.LOST_STRUCTURE]


def test_a_cdata_section_handed_over_as_a_comment_is_read_as_text() -> None:
    """Later Pythons hand a CDATA section outside foreign content to `handle_comment`; the reader
    takes it from there too, and an ordinary comment is still nothing. Delete this and code and
    link text vanish from every storage-format page on the Python after this one."""
    reader = sop_files._PageReader(export=False)
    reader.feed("<p>")
    reader.handle_comment("[CDATA[make account]]")
    reader.handle_comment(" an ordinary comment ")
    reader.feed("</p>")
    reader.close()

    assert reader.writer.text() == "make account"


# ---------------------------------------------------------------------- refusals
def a_small_word() -> bytes:
    return word(a_titled)


@pytest.mark.parametrize(
    ("file_name", "content", "refusal"),
    [
        ("Procedure.pdf", b"%PDF-1.7", "neither"),
        ("Procedure.doc", b"anything", "save it as a .docx"),
        ("Procedure.docx", b"", "the file is empty"),
        ("Procedure.docx", b"<html><p>Steps</p></html>", "not a Word document"),
        ("Procedure.html", b"PK\x03\x04 a zip", "its bytes are not text"),
        ("Procedure.html", "Étapes".encode("latin-1") + b" " * 20, "not UTF-8"),
        ("Procedure.html", b"<p>Steps" + bytes(1) + b"</p>" + b" " * 20, "its bytes are not text"),
    ],
    ids=["pdf", "old word", "empty", "page named docx", "zip named html", "latin-1", "nul"],
)
def test_a_file_that_is_not_what_its_name_says_is_refused_before_it_is_parsed(
    file_name: str, content: bytes, refusal: str
) -> None:
    """The name picks the reader and the bytes must agree with it, or nothing reads them. Delete
    this and a zip named as a page, or a page named as a Word document, reaches the parser the
    name picked."""
    with pytest.raises(SopError, match=re.escape(refusal)):
        procedure_text(file_name, content)


def test_a_zip_that_is_not_a_word_document_is_refused_and_a_word_document_is_read() -> None:
    """A `.docx` that is a zip of something else is refused in words; the real one beside it is
    read. Delete this and a zip of anything reaches `python-docx`, or every Word file is refused
    and every refusal test above still passes."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("notes.txt", "Steps")

    with pytest.raises(SopError, match="not a Word document"):
        procedure_text("Procedure.docx", out.getvalue())
    assert procedure_text("Procedure.docx", a_small_word()) == ProcedureText(
        text="# Raising an invoice", source=SourceFormat.WORD, fallback_name="Procedure"
    )


def test_a_file_over_the_bound_is_refused_for_its_size_and_one_under_it_is_read() -> None:
    """Refused on its length alone: the file past the bound is a real Word document. Delete this
    and a procedure can be any size the web server accepts, held whole and then parsed."""
    padded = with_part(
        a_small_word(), "word/media/padding.bin", bytes(MAX_PROCEDURE_BYTES), stored=True
    )

    assert len(padded) > MAX_PROCEDURE_BYTES
    with pytest.raises(SopError, match=f"over the {MAX_PROCEDURE_BYTES} a procedure may be"):
        procedure_text("Procedure.docx", padded)
    assert procedure_text("Procedure.docx", a_small_word()).text == "# Raising an invoice"


@pytest.mark.parametrize(
    ("extra", "refusal"),
    [
        (
            lambda content: with_part(
                content, "word/media/big.bin", bytes(MAX_PROCEDURE_EXPANDED_BYTES)
            ),
            f"over the {MAX_PROCEDURE_EXPANDED_BYTES} a procedure may",
        ),
        (
            lambda content: _many_parts(content, MAX_PROCEDURE_PARTS),
            f"over the {MAX_PROCEDURE_PARTS} a procedure may hold",
        ),
        (lambda content: with_part(content, "word/vbaProject.bin", b"macros"), "carries macros"),
        (
            lambda content: with_part(
                content, "word/extra.xml", b"<!DOCTYPE x [<!ENTITY a 'b'>]><x>&a;</x>"
            ),
            "declares XML entities",
        ),
    ],
    ids=["expands past the bound", "too many parts", "macros", "entities"],
)
def test_a_word_document_is_judged_as_an_archive_before_it_is_opened(
    extra: Callable[[bytes], bytes], refusal: str
) -> None:
    """This module's bounds on the parts and their unpacked size, then the structural check's
    macros and entities. Delete this and a small file that inflates to gigabytes, or one carrying
    a macro project or an entity declaration, is handed to `python-docx`."""
    with pytest.raises(SopError, match=re.escape(refusal)):
        procedure_text("Procedure.docx", extra(a_small_word()))


def _many_parts(content: bytes, bound: int) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(out, "w") as copy:
        for one in source.infolist():
            copy.writestr(one, source.read(one))
        for index in range(bound):
            copy.writestr(f"word/media/{index}.bin", b"")
    return out.getvalue()


# ---------------------------------------------------------------------- names
def test_the_format_is_the_name_s_suffix_and_nothing_else() -> None:
    """Delete this and a `.docm` carrying macros, or a `.txt`, is read as one of the two."""
    assert format_of("Invoice.DOCX") is SourceFormat.WORD
    assert [format_of(f"page{suffix}") for suffix in (".html", ".htm", ".xhtml", ".xml")] == [
        SourceFormat.CONFLUENCE
    ] * 4
    assert [format_of(name) for name in ("x.docm", "x.doc", "x.txt", "docx")] == [None] * 4


@pytest.mark.parametrize(
    ("given", "kept"),
    [
        ("Raising an invoice (v2).docx", "Raising-an-invoice-v2.docx"),
        ("C:" + chr(92) + "Users" + chr(92) + "me" + chr(92) + "SOP.DOCX", "SOP.docx"),
        ("folder/..hidden page.HTML", "hidden-page.html"),
        ("().docx", "procedure.docx"),
    ],
)
def test_a_file_name_is_kept_as_its_last_segment_in_the_member_grammar(
    given: str, kept: str
) -> None:
    """Spaces and brackets become hyphens, a path is cut to its last segment, and the suffix is
    kept lower case, so the upload's location is one the skill source accepts. Delete this and
    nearly every real Word file name is refused by `SkillSource`, or a path is stored."""
    assert kept_name(given) == kept
    assert safe_archive_member(kept_name(given)) == kept
