"""Small real documents, built in memory: a PDF with a text layer, a Word file, and hostile zips.

Built rather than committed, so each test says what is in its file and a binary nobody can read in
review never sits in the repository. The PDF is written by hand, object by object with its offsets
counted, because the product depends on a PDF reader and not on a PDF writer; the Word file is
written by python-docx, which the product already depends on to read one.

Task ids: none
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Sequence

import docx
from pypdf import PdfReader, PdfWriter

#: What separates two lines inside a string built here.
LINE = chr(10)


def a_pdf(*pages: str, extra: str = "") -> bytes:
    """A PDF with one line of Helvetica text per page, and an optional extra catalog entry.

    The text must not contain a parenthesis or a backslash, which a PDF string would need
    escaped. `extra` is written into the catalog dictionary as given, which is how a test puts an
    action in a file.
    """
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append(f"<< /Type /Catalog /Pages 2 0 R {extra}>>".encode())
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font = 3 + 2 * len(pages)
    for index, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            (
                "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {4 + 2 * index} 0 R >>"
            ).encode()
        )
        objects.append(
            f"<< /Length {len(content)} >>{LINE}stream{LINE}".encode()
            + content
            + f"{LINE}endstream".encode()
        )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(f"%PDF-1.4{LINE}".encode())
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj{LINE}".encode() + body + f"{LINE}endobj{LINE}".encode())
    xref = out.tell()
    out.write(f"xref{LINE}0 {len(objects) + 1}{LINE}".encode())
    out.write(f"0000000000 65535 f {LINE}".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n {LINE}".encode())
    out.write(
        (
            f"trailer{LINE}<< /Size {len(objects) + 1} /Root 1 0 R >>{LINE}"
            f"startxref{LINE}{xref}{LINE}%%EOF{LINE}"
        ).encode()
    )
    return out.getvalue()


#: The word a locked PDF opens with. Invented here, and a test file's, never anybody's.
OPENING_WORD = "tealopening"


def a_locked_pdf(*pages: str) -> bytes:
    """The same PDF, encrypted with a word needed to open it."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(a_pdf(*pages))))
    writer.encrypt(user_password=OPENING_WORD, algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def an_owner_locked_pdf(*pages: str) -> bytes:
    """A PDF with restrictions and no password to open it, which is most "protected" PDFs."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(a_pdf(*pages))))
    writer.encrypt(user_password="", owner_password="owner", algorithm="RC4-128")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def a_word_document(
    *,
    heading: str = "",
    paragraphs: Sequence[str] = (),
    table: Sequence[Sequence[str]] = (),
    after: Sequence[str] = (),
) -> bytes:
    """A Word file: an optional heading, paragraphs, an optional table, then more paragraphs."""
    document = docx.Document()
    if heading:
        document.add_heading(heading, level=1)
    for one in paragraphs:
        document.add_paragraph(one)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, cells in zip(grid.rows, table, strict=True):
            for cell, text in zip(row.cells, cells, strict=True):
                cell.text = text
    for one in after:
        document.add_paragraph(one)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


#: The signatures of a zip's local and central headers, and where each keeps its flag word.
_LOCAL_HEADER: bytes = bytes((0x50, 0x4B, 0x03, 0x04))
_CENTRAL_HEADER: bytes = bytes((0x50, 0x4B, 0x01, 0x02))
_LOCAL_FLAG_AT = 6
_CENTRAL_FLAG_AT = 8


def a_zip(parts: dict[str, bytes], *, encrypted: bool = False) -> bytes:
    """A zip holding these parts, deflated, and every part flagged as encrypted if asked.

    The flag is set in the written bytes because `zipfile` cannot write an encrypted part and
    clears the bit when asked to; a scanner reads the flag, which is the claim being tested.
    """
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    written = bytearray(out.getvalue())
    if encrypted:
        for signature, at in ((_LOCAL_HEADER, _LOCAL_FLAG_AT), (_CENTRAL_HEADER, _CENTRAL_FLAG_AT)):
            start = written.find(signature)
            while start != -1:
                written[start + at] |= 0x1
                start = written.find(signature, start + 1)
    return bytes(written)


def with_part(document: bytes, name: str, body: bytes) -> bytes:
    """A copy of an Office file with one more part added."""
    source = zipfile.ZipFile(io.BytesIO(document))
    parts = {one.filename: source.read(one) for one in source.infolist()}
    parts[name] = body
    return a_zip(parts)
