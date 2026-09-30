"""Opening the file a written procedure was saved in: a Word document or a Confluence page.

`brain.tools.sop_import` reads a procedure's text and decides what a reviewer is shown about it,
and said in its own docstring that nothing opened the file. This is that half: it takes the bytes
an administrator chose and the name they carried, refuses anything that is not what it says it is
before a parser sees it, and writes the procedure out as Markdown with its structure kept. Nothing
here decides what the text means; every judgement about the words is `sop_import`'s.

**The name declares the format and the bytes must agree, before anything parses.** A `.docx` is a
zip and a Confluence export is text, and `brain.knowledge.ingest.sniff` says which the bytes are.
A file named for one and built as the other is refused, because the parser the name picks is the
one it would attack. The format is taken from the name rather than sniffed alone for
`brain.tools.sop_import.SourceFormat`'s reason: which cleanups run is the caller's declaration.
See `A_FILE_IS_JUDGED_BEFORE_IT_IS_PARSED`.

**A Word document is judged as the archive it is, twice.** First against this module's own
bounds, which are a procedure's and much tighter than a knowledge upload's: a few hundred parts and
tens of megabytes unpacked. Then by `brain.knowledge.text_path.StructuralCheck`, reused rather
than restated, which refuses an entry that expands past any document's ratio, a macro project, an
encrypted part and an XML part declaring entities, because `python-docx` is the parser those
shapes attack there and here alike. Rejected: `brain.tools.extract`'s rules alone. They were
written for a skill archive of text, and refuse the images an ordinary procedure carries.

**A Confluence page is read by the standard library's HTML parser, and nothing new was added.**
The export a person saves from a page and the storage format a page is kept in are both markup
that `html.parser` reads without building a tree or expanding an entity, so a page cannot make
this process fetch, include or inflate anything. Rejected: `lxml.html`, installed only because
`python-docx` needs it, whose recovering parser is a larger surface for no gain on these files;
and BeautifulSoup, a new dependency for what forty lines of state do here.

**Structure is written as Markdown, in one shape whatever the document used.** A heading is a
heading because its style says so (Title, Heading 1 to 9, an outline level) or because the whole
paragraph is bold and short, which is how people who never learned styles write one; it is written
with hashes, so `sop_import` is told the headings are marked and reads a numbered line as a step.
A step is numbered in order whether Word numbered it, the author typed `1)`, `2.` or `Step 3:`, or
it sat in a table whose first column counts. See `A_STEP_IS_NUMBERED_IN_ORDER_WHATEVER_ITS_STYLE`.

**What cannot be carried is written into the text as a note, and never dropped.** A tracked change
nobody accepted, a comment, text a reader cannot see, a macro whose content is not in the export
and a list whose numbering the file does not define each leave a bracketed note on its line, from
the vocabulary `sop_import` reads. The notes are in the draft's body, so the findings the importer
is shown and the findings a reviewer is shown later are one list read from one text. See
`A_NOTE_IN_THE_TEXT_IS_A_FINDING_EVERY_REVIEWER_SEES`.

**Pictures are not carried and are not flagged.** A skill is text a model reads, so a screenshot
has nowhere to go; a note beside each would bury the one finding that matters in a procedure with
thirty of them. Said here rather than on every draft.

Nothing here opens a socket, reads a clock or writes a file.

Task ids: M12.2.10
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any, Final

import docx
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from brain.knowledge.ingest import SNIFF_BYTES, Container, ScanCause, ScanVerdict, sniff
from brain.knowledge.text_path import StructuralCheck, pipe_row
from brain.tools.extract import MAX_MEMBERS
from brain.tools.sop_import import (
    COMMENT_NOTE,
    HIDDEN_TEXT_NOTE,
    NOT_READ_NOTE,
    TRACKED_CHANGE_NOTE,
    SopError,
    SourceFormat,
)

# ------------------------------------------------------------------ written-down reasons
#: Why a file is refused before any parser reads it.
A_FILE_IS_JUDGED_BEFORE_IT_IS_PARSED: Final = (
    "the parser a file's name picks is the parser a hostile file attacks, so the name, the size, "
    "the container the bytes prove and, for a Word document, the archive's parts are all judged "
    "before python-docx or the HTML reader sees a byte"
)

#: Why every numbering style becomes one.
A_STEP_IS_NUMBERED_IN_ORDER_WHATEVER_ITS_STYLE: Final = (
    "a procedure's steps arrive numbered by Word, typed by hand as 1) or Step 3:, or counted in "
    "a table's first column, and a reader of the draft needs one list in the document's order. "
    "The numbers restart under each heading and at each new Word list, so a step's number is "
    "its place in its own section"
)

#: Why what could not be carried is written into the text.
A_NOTE_IN_THE_TEXT_IS_A_FINDING_EVERY_REVIEWER_SEES: Final = (
    "a finding kept only in the answer to the import is shown to the importer and to nobody who "
    "reviews the draft afterwards. Written into the body as a bracketed note on its line, it is "
    "read back as the same finding by whoever opens the draft, and deleting the note is an edit "
    "the review shows"
)

# ------------------------------------------------------------------------ the figures
#: The one name a Word document is imported under. `.doc` is Word's older binary format, which
#: no parser here reads; `.docm` carries macros, which a procedure has no use for.
WORD_SUFFIX: Final = ".docx"

#: The names a Confluence page is imported under: its HTML export, and the storage format a page
#: is kept in, which is XHTML and read by the same reader.
PAGE_SUFFIXES: Final[tuple[str, ...]] = (".html", ".htm", ".xhtml", ".xml")

#: The largest file a procedure may be. A written procedure with screenshots is a few megabytes;
#: ten is room for a long one, and past it the file is something else.
MAX_PROCEDURE_BYTES: Final = 10 * 1024 * 1024

#: The most parts a Word document may hold: `brain.tools.extract`'s bound for an archive. A
#: procedure has a few dozen; hundreds is a construction.
MAX_PROCEDURE_PARTS: Final = MAX_MEMBERS

#: The most a Word document's parts may add up to unpacked. Tighter than the knowledge path's
#: bound, because every part is read into memory when the document opens.
MAX_PROCEDURE_EXPANDED_BYTES: Final = 64 * 1024 * 1024

#: The longest paragraph read as a heading because it is bold. A bold sentence is emphasis.
MAX_BOLD_HEADING_CHARS: Final = 120

#: How far up a style's parents a heading level or a list is looked for. Word's own chains are
#: two or three deep; the bound stops a style that names itself as its own parent.
STYLE_DEPTH: Final = 10

#: What indents a list item one level, which lines up under the text of `1. `.
INDENT: Final = "   "

#: A step typed by hand: `1.`, `2)`, `(3)` or `Step 4:`, then the step. The punctuation is
#: required and followed by space, so a numbered heading like `1.2 Scope` is not a step.
TYPED_STEP_RE: Final = re.compile(r"^\s*(?:step\s*)?\(?(\d{1,3})\s*[.):]\s+(\S.*)$", re.IGNORECASE)

#: A cell that only counts: `1`, `2.` or `3)`.
COUNT_CELL_RE: Final = re.compile(r"^\d{1,3}[.)]?$")

#: The first cell of a table whose rows are steps.
STEP_COLUMN_WORDS: Final[frozenset[str]] = frozenset(
    {"step", "steps", "no", "no.", "#", "number", "seq", "order", "step no", "step no."}
)

#: Word's built-in heading styles by name, which every language pack keeps underneath.
HEADING_STYLE_RE: Final = re.compile(r"^heading ([1-9])$")

#: Marks a Confluence export: the page's own content sits inside this element, and everything
#: outside it (breadcrumbs, attachments, the footer) is the export's frame.
MAIN_CONTENT_RE: Final = re.compile(r"""\bid\s*=\s*["']main-content["']""")

#: `Space : Page`, the title a Confluence export gives a page.
SPACE_PREFIX_RE: Final = re.compile(r"^[^:]{1,120} : (?=\S)")

#: Elements whose contents are never words of the procedure: code, styling, and what a Confluence
#: macro is configured with rather than what it shows.
SKIPPED_TAGS: Final[frozenset[str]] = frozenset(
    {
        "script",
        "style",
        "head",
        "noscript",
        "template",
        "svg",
        "button",
        "nav",
        "ac:parameter",
        "ac:task-id",
        "ac:task-status",
        "ac:task-uuid",
        "ac:emoticon",
        "ac:image",
        "ac:placeholder",
    }
)

#: A Confluence export's own frame, by class, when it sits inside the content.
SKIPPED_CLASSES: Final[frozenset[str]] = frozenset({"expand-control", "page-metadata"})

#: Elements that end a line of text.
BLOCK_TAGS: Final[frozenset[str]] = frozenset(
    {
        "p",
        "div",
        "br",
        "hr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "ol",
        "ul",
        "table",
        "tr",
        "td",
        "th",
        "pre",
        "blockquote",
        "section",
        "article",
        "header",
        "footer",
        "dl",
        "dt",
        "dd",
        "figure",
        "figcaption",
        "caption",
        "ac:structured-macro",
        "ac:rich-text-body",
        "ac:task-list",
        "ac:task",
        "ac:task-body",
        "ac:layout",
        "ac:layout-section",
        "ac:layout-cell",
    }
)

#: Elements that never have an end tag.
VOID_TAGS: Final[frozenset[str]] = frozenset(
    {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "source", "wbr"}
)

#: Confluence's panels, by macro name and by the class its export gives them, and the word each
#: is labelled with in the draft.
PANEL_LABELS: Final[dict[str, str]] = {
    "info": "Info",
    "information": "Info",
    "note": "Note",
    "warning": "Warning",
    "tip": "Tip",
    "panel": "",
}

#: Macros whose body is code, written as a fenced block.
CODE_MACROS: Final[frozenset[str]] = frozenset({"code", "noformat"})

#: Macros that show navigation rather than content, left out without a note.
NAVIGATION_MACROS: Final[frozenset[str]] = frozenset(
    {"toc", "toc-zone", "anchor", "children", "pagetree", "recently-updated"}
)

#: What the importer is told for each shape the structural check refuses a Word document for,
#: in a procedure's words rather than a knowledge upload's.
ARCHIVE_REFUSAL: Final[dict[ScanCause, str]] = {
    ScanCause.MACROS: (
        "it carries macros, which a procedure has no use for; save it again from Word as a .docx"
    ),
    ScanCause.EXPANDS_WITHOUT_BOUND: (
        "it unpacks to far more than any document holds, which is how a file is built to exhaust "
        "whatever reads it; save it again from Word"
    ),
    ScanCause.XML_ENTITIES: (
        "one of its parts declares XML entities, which a Word document never does; save it "
        "again from Word"
    ),
    ScanCause.ENCRYPTED_PART: "part of it is encrypted; save a copy with the password removed",
    ScanCause.UNOPENABLE: "it could not be opened as a zip, which a .docx is; save it again",
}


# ------------------------------------------------------------------------ the answer
@dataclass(frozen=True)
class ProcedureText:
    """A procedure's text as Markdown, the format it was read as, and a name to fall back on.

    `fallback_name` is the file's own name without its suffix, used only when the document has no
    title a skill could be called by.
    """

    text: str
    source: SourceFormat
    fallback_name: str


def _refused(reason: str) -> SopError:
    return SopError(f"this procedure was not imported: {reason}")


def format_of(file_name: str) -> SourceFormat | None:
    """The format a file's name declares, or None for a name no procedure is imported under."""
    lowered = file_name.strip().lower()
    if lowered.endswith(WORD_SUFFIX):
        return SourceFormat.WORD
    if lowered.endswith(PAGE_SUFFIXES):
        return SourceFormat.CONFLUENCE
    return None


def kept_name(file_name: str) -> str:
    """The file's name as the library keeps it: its last segment, in the archive-member grammar.

    `brain.tools.skills.SkillSource` holds an upload's location to letters, digits, dots, hyphens
    and underscores, and a Word document is named the way people name documents, with spaces and
    brackets. Refusing those would refuse nearly every real procedure, so each run of anything
    else becomes a hyphen and the suffix is kept, which is how the format is recorded.
    """
    base = file_name.replace(chr(92), "/").rsplit("/", 1)[-1].strip()
    stem, _dot, suffix = base.rpartition(".")
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-._")[:80] or "procedure"
    return f"{cleaned}.{suffix.lower()}"


def procedure_text(file_name: str, content: bytes) -> ProcedureText:
    """The procedure in a Word document or a Confluence page, or a refusal before it is parsed.

    See `A_FILE_IS_JUDGED_BEFORE_IT_IS_PARSED`. Every refusal is `SopError`, in words the person
    who chose the file can act on, since they already hold it.
    """
    source = format_of(file_name)
    if source is None:
        if file_name.strip().lower().endswith(".doc"):
            raise _refused(
                "this is Word's older .doc format; open it in Word and save it as a .docx"
            )
        raise _refused(
            "a procedure is imported from a Word document saved as .docx, or from a Confluence "
            "page exported as .html; this file is neither"
        )
    if not content:
        raise _refused("the file is empty")
    if len(content) > MAX_PROCEDURE_BYTES:
        raise _refused(
            f"the file is {len(content)} bytes, over the {MAX_PROCEDURE_BYTES} a procedure may be"
        )
    container = sniff(content[:SNIFF_BYTES])
    fallback = kept_name(file_name).rsplit(".", 1)[0]
    if source is SourceFormat.WORD:
        if container is not Container.ZIP:
            raise _refused(
                "the file is named .docx and is not a Word document, whose bytes are a zip; "
                "save it again from Word as a .docx"
            )
        _judge_archive(content)
        return ProcedureText(text=_word_text(content), source=source, fallback_name=fallback)
    if container is not Container.TEXT:
        raise _refused(
            "the file is named as a web page and its bytes are not text; export the page from "
            "Confluence as HTML and import that"
        )
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise _refused("the page is not UTF-8 text, which a Confluence export is") from None
    if chr(0) in decoded:
        raise _refused("the page holds a byte no text contains")
    return ProcedureText(text=_page_text(decoded), source=source, fallback_name=fallback)


def _judge_archive(content: bytes) -> None:
    """This module's bounds on a Word document's parts, then the structural check's verdict."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            parts = archive.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError, EOFError):
        raise _refused(ARCHIVE_REFUSAL[ScanCause.UNOPENABLE]) from None
    if len(parts) > MAX_PROCEDURE_PARTS:
        raise _refused(
            f"the document holds {len(parts)} parts, over the {MAX_PROCEDURE_PARTS} a procedure "
            "may hold"
        )
    unpacked = sum(one.file_size for one in parts)
    if unpacked > MAX_PROCEDURE_EXPANDED_BYTES:
        raise _refused(
            f"the document unpacks to {unpacked} bytes, over the {MAX_PROCEDURE_EXPANDED_BYTES} "
            "a procedure may"
        )
    report = StructuralCheck().scan(content)
    if report.verdict is not ScanVerdict.CLEAN:
        cause = report.cause or ScanCause.UNSTATED
        raise _refused(
            ARCHIVE_REFUSAL.get(cause, "the document could not be checked; save it again")
        )


# ------------------------------------------------------------------------ writing Markdown
@dataclass
class _Writer:
    """Lines of Markdown, one shape for both readers, and the step counters per list level.

    Consecutive list items and consecutive quoted lines sit together; every other block is set off
    by a blank line. See `A_STEP_IS_NUMBERED_IN_ORDER_WHATEVER_ITS_STYLE` for when counting starts
    again.
    """

    lines: list[str] = field(default_factory=list)
    last: str = ""
    counters: list[int] = field(default_factory=list)
    #: The Word list, or `typed`, the counters at the top level belong to.
    list_id: str | None = None
    #: How many panels deep the next line sits, each one a `> `.
    quote_depth: int = 0

    def _put(self, line: str, kind: str) -> None:
        quoted = self.quote_depth > 0
        prefix = "> " * self.quote_depth
        kind = f"quote-{kind}" if quoted else kind
        together = kind == self.last and kind.endswith("list")
        if self.lines and not together:
            # Inside a panel the gap is a quoted empty line, so the panel stays one quote.
            inside = quoted and self.last.startswith("quote")
            self.lines.append(prefix.rstrip() if inside else "")
        self.lines.append(chr(10).join(prefix + one for one in line.split(chr(10))))
        self.last = kind

    def restart(self) -> None:
        self.counters.clear()
        self.list_id = None

    def restart_level(self, level: int) -> None:
        del self.counters[level:]

    def heading(self, level: int, text: str) -> None:
        self.restart()
        self._put(f"{'#' * max(1, min(6, level))} {text}", "heading")

    def paragraph(self, text: str) -> None:
        typed = TYPED_STEP_RE.match(text)
        if typed:
            # A typed 1 starts a list; any other typed number carries on the one being counted,
            # which is what somebody typing 3) after two numbered steps meant.
            if int(typed.group(1)) <= 1:
                self.counters.clear()
                self.list_id = "typed"
            self.numbered(0, typed.group(2).strip())
            return
        # A paragraph opening with a hash would be read back as a heading it never was.
        self._put(chr(92) + text if text.startswith("#") else text, "para")

    def numbered(self, level: int, text: str, *, list_id: str | None = None) -> None:
        if level == 0 and list_id is not None:
            if self.list_id is not None and list_id != self.list_id:
                self.counters.clear()
            self.list_id = list_id
        del self.counters[level + 1 :]
        while len(self.counters) <= level:
            self.counters.append(0)
        self.counters[level] += 1
        self._put(f"{INDENT * level}{self.counters[level]}. {text}", "list")

    def bullet(self, level: int, text: str) -> None:
        del self.counters[level + 1 :]
        self._put(f"{INDENT * level}- {text}", "list")

    def continued(self, level: int, text: str) -> None:
        """A second paragraph inside a list item, indented under it."""
        self._put(f"{INDENT * (level + 1)}{text}", "list")

    def note(self, what: str) -> None:
        self._put(f"{NOT_READ_NOTE}{what}]", "para")

    def code(self, text: str) -> None:
        fence = "```"
        self._put(f"{fence}{chr(10)}{text}{chr(10)}{fence}", "code")

    def table(self, rows: Sequence[Sequence[str]]) -> None:
        """A table of steps as numbered steps, and any other table as a pipe table.

        A steps table whose count starts past one carries on the list being counted, as a typed
        `3)` does; one counting from one, or not counting at all, is a list of its own.
        """
        kept = [list(row) for row in rows if any(cell.strip() for cell in row)]
        if not kept:
            return
        header, body, counted = _steps_of(kept)
        if header is None and not counted:
            width = max(len(row) for row in kept)
            padded = [row + [""] * (width - len(row)) for row in kept]
            rule = "| " + " | ".join("---" for _ in range(width)) + " |"
            lines = [pipe_row(padded[0]), rule, *(pipe_row(row) for row in padded[1:])]
            self._put(chr(10).join(lines), "table")
            return
        opening = body[0][0].strip().rstrip(".)") if counted and body and body[0] else ""
        if not (opening.isdigit() and int(opening) > 1):
            self.counters.clear()
            self.list_id = "table"
        for row in body:
            cells = row[1:] if counted else row
            step = _step_line(cells, header[1:] if (header and counted) else header)
            if step:
                self.numbered(0, step)

    def text(self) -> str:
        return chr(10).join(self.lines).strip()


def _folded(cell: str) -> str:
    return " ".join(cell.lower().split())


def _steps_of(
    rows: list[list[str]],
) -> tuple[list[str] | None, list[list[str]], bool]:
    """The header, the step rows and whether the first column counts, for a table of steps.

    A table is one of steps when its first header cell says so (`Step`, `No.`, `#`), or when every
    row under a first row that does not count opens with a count. Anything else is a table, and
    the empty header with no counting column says so.
    """
    first = rows[0]
    rest = rows[1:]
    counted_rest = bool(rest) and all(row and COUNT_CELL_RE.match(row[0].strip()) for row in rest)
    if _folded(first[0]) in STEP_COLUMN_WORDS if first else False:
        return first, rest, counted_rest
    if counted_rest and not COUNT_CELL_RE.match(first[0].strip()):
        return first, rest, True
    if len(rows) > 1 and all(row and COUNT_CELL_RE.match(row[0].strip()) for row in rows):
        return None, rows, True
    return None, [], False


def _step_line(cells: Sequence[str], header: Sequence[str] | None) -> str:
    """One row of a steps table as one step: its first filled cell, then the rest by heading."""
    filled = [(index, cell.strip()) for index, cell in enumerate(cells) if cell.strip()]
    if not filled:
        return ""
    (_first, step), extras = filled[0], filled[1:]
    if not extras:
        return step
    labelled = []
    for index, cell in extras:
        label = header[index].strip() if header is not None and index < len(header) else ""
        labelled.append(f"{label}: {cell}" if label else cell)
    return f"{step} ({'; '.join(labelled)})"


# ------------------------------------------------------------------------ Word
def _xpath(element: Any, path: str) -> list[Any]:
    """The nodes or strings a query from one of `python-docx`'s own elements finds.

    Its elements type `xpath` as Any and map the `w` prefix for the query, which a plain element
    of the file (a numbering definition, a content control) does not: those are walked with
    `_w` and the tree's own `find`, never queried here.
    """
    found = element.xpath(path)
    return list(found) if isinstance(found, list) else []


def _w(name: str) -> str:
    """A WordprocessingML name as the tree spells it, for `find`, `findall` and `get`."""
    return str(qn(f"w:{name}"))


def _numbering_formats(document: WordDocument) -> dict[tuple[str, str], str]:
    """Every list definition's number format, by the list's id and the level, from the file.

    A level whose format is not stated is decimal, as Word reads it; a list the file names and
    does not define is absent here, which is how an undefined list is told apart.
    """
    try:
        numbering = document.part.numbering_part.element
    except (KeyError, NotImplementedError, ValueError):
        return {}
    abstract: dict[str, dict[str, str]] = {}
    for one in numbering.findall(_w("abstractNum")):
        abstract[one.get(_w("abstractNumId")) or ""] = _levels(one.findall(_w("lvl")))
    formats: dict[tuple[str, str], str] = {}
    for num in numbering.findall(_w("num")):
        pointer = num.find(_w("abstractNumId"))
        shared = "" if pointer is None else pointer.get(_w("val")) or ""
        levels = abstract.get(shared, {}) | _levels(num.findall(f"{_w('lvlOverride')}/{_w('lvl')}"))
        for ilvl, fmt in levels.items():
            formats[(num.get(_w("numId")) or "", ilvl)] = fmt
    return formats


def _levels(levels: Iterable[Any]) -> dict[str, str]:
    """Each level's number format by its index; a level stating none is decimal, as in Word."""
    found: dict[str, str] = {}
    for level in levels:
        fmt = level.find(_w("numFmt"))
        stated = None if fmt is None else fmt.get(_w("val"))
        found[level.get(_w("ilvl")) or "0"] = stated or "decimal"
    return found


def _styles(paragraph: Paragraph) -> Iterable[Any]:
    """The paragraph's style and its parents, nearest first, at most `STYLE_DEPTH`."""
    style: Any = paragraph.style
    for _ in range(STYLE_DEPTH):
        if style is None:
            return
        yield style
        style = style.base_style


def _heading_level(paragraph: Paragraph, element: Any) -> int | None:
    """The Markdown level of a heading paragraph, by its style or its outline level, or None."""
    outline = _xpath(element, "./w:pPr/w:outlineLvl/@w:val")
    for style in _styles(paragraph):
        name = (style.name or "").strip().lower()
        if name == "title":
            return 1
        found = HEADING_STYLE_RE.match(name)
        if found:
            return min(6, int(found.group(1)) + 1)
        outline = outline or _xpath(style.element, "./w:pPr/w:outlineLvl/@w:val")
    if outline and outline[0].isdigit() and int(outline[0]) < 9:
        return min(6, int(outline[0]) + 2)
    return None


def _is_contents(paragraph: Paragraph) -> bool:
    """A table of contents' line, which is a copy of the headings and not a word of the text."""
    return any((style.name or "").strip().lower().startswith("toc") for style in _styles(paragraph))


def _bold_heading(paragraph: Paragraph, text: str) -> bool:
    """A short paragraph set wholly in bold: the heading of somebody who never used styles."""
    if len(text) > MAX_BOLD_HEADING_CHARS or text.endswith((".", "!", "?", ",", ";")):
        return False
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False
    style_bold = any(style.font.bold is True for style in _styles(paragraph))
    return all(run.bold is True or (run.bold is None and style_bold) for run in runs)


def _list_of(
    paragraph: Paragraph, element: Any, formats: dict[tuple[str, str], str]
) -> tuple[str, int, str] | None:
    """`(kind, level, list id)` for a list paragraph, or None. Kind is numbered, bullet or lost.

    The numbering is the paragraph's own, or its style's, as Word resolves it. A list id the file
    does not define is `lost`: the item is kept, and its place in the steps is marked unknown.
    """
    found = _xpath(element, "./w:pPr/w:numPr")
    if not found:
        for style in _styles(paragraph):
            found = _xpath(style.element, "./w:pPr/w:numPr")
            if found:
                break
    if not found:
        return None
    num_id = "".join(_xpath(found[0], "./w:numId/@w:val"))
    ilvl = "".join(_xpath(found[0], "./w:ilvl/@w:val")) or "0"
    if num_id in ("", "0"):
        # Numbering switched off on a paragraph whose style would number it.
        return None
    level = min(int(ilvl), 8) if ilvl.isdigit() else 0
    fmt = formats.get((num_id, ilvl))
    if fmt is None:
        return "lost", level, num_id
    if fmt == "none":
        return None
    return ("bullet" if fmt == "bullet" else "numbered"), level, num_id


#: A run a reader cannot see: hidden, or white.
_HIDDEN_RUNS: Final = (
    ".//w:r[w:rPr/w:vanish[not(@w:val) or @w:val='1' or @w:val='true' or @w:val='on']]"
    " | .//w:r[translate(w:rPr/w:color/@w:val, 'abcdef', 'ABCDEF') = 'FFFFFF']"
)


def _paragraph_text(paragraph: Paragraph, element: Any) -> str:
    """The paragraph's words as Word's own reader gives them, with a note for what it holds unseen.

    `python-docx` reads the runs a paragraph holds directly, so a tracked insertion and a tracked
    deletion are both left out of the words; the note says the paragraph has them.
    """
    text = " ".join(paragraph.text.split())
    notes = []
    if _xpath(element, ".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo"):
        notes.append(TRACKED_CHANGE_NOTE)
    if _xpath(element, ".//w:commentReference | .//w:commentRangeStart"):
        notes.append(COMMENT_NOTE)
    if _xpath(element, _HIDDEN_RUNS):
        notes.append(HIDDEN_TEXT_NOTE)
    return " ".join([text, *notes]).strip() if notes else text


def _blocks(element: Any) -> Iterable[Any]:
    """The paragraphs and tables under a body, in order, through content controls and custom XML.

    `python-docx` lists only the paragraphs and tables directly under the body, so a section a
    template wraps in a content control would be skipped whole and nobody told.
    """
    for child in element.iterchildren():
        if child.tag in (_w("p"), _w("tbl")):
            yield child
        elif child.tag == _w("sdt"):
            for content in child.findall(_w("sdtContent")):
                yield from _blocks(content)
        elif child.tag == _w("customXml"):
            yield from _blocks(child)


def _word_text(content: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(content))
    except (PackageNotFoundError, KeyError, ValueError):
        raise _refused(
            "the file is a zip and not a Word document; save it again from Word as a .docx"
        ) from None
    except zipfile.BadZipFile:
        raise _refused(ARCHIVE_REFUSAL[ScanCause.UNOPENABLE]) from None
    formats = _numbering_formats(document)
    writer = _Writer()
    for element in _blocks(document.element.body):
        if element.tag == _w("tbl"):
            writer.table(_word_rows(document, element))
            continue
        paragraph = Paragraph(element, document)
        if _is_contents(paragraph):
            continue
        text = _paragraph_text(paragraph, element)
        if not text:
            continue
        level = _heading_level(paragraph, element)
        listed = _list_of(paragraph, element, formats)
        if level is not None:
            writer.heading(level, text)
        elif listed is not None:
            kind, depth, num_id = listed
            if kind == "numbered":
                writer.numbered(depth, text, list_id=num_id)
            elif kind == "bullet":
                writer.bullet(depth, text)
            else:
                writer.bullet(depth, f"{text} {NOT_READ_NOTE}the number of this step]")
        elif _bold_heading(paragraph, text):
            writer.heading(2, text.rstrip(":").strip())
        else:
            writer.paragraph(text)
    return writer.text()


def _word_rows(document: WordDocument, table: Any) -> list[list[str]]:
    """A table's rows as the text of each cell, one cell per cell the file holds.

    Read from the cells the file lists rather than from `python-docx`'s grid, which repeats a
    merged cell once for every column it spans and would write its words into the step twice.
    """
    rows: list[list[str]] = []
    for row in _xpath(table, "./w:tr"):
        cells = []
        for cell in _xpath(row, "./w:tc | ./w:sdt/w:sdtContent/w:tc"):
            words = [
                _paragraph_text(Paragraph(one, document), one) for one in _xpath(cell, "./w:p")
            ]
            cells.append(" ".join(word for word in words if word))
        rows.append(cells)
    return rows


# ------------------------------------------------------------------------ Confluence
@dataclass
class _Open:
    """One element the page reader has open, and what it changed that its end has to undo."""

    tag: str
    effects: list[str] = field(default_factory=list)
    macro: str = ""
    had_body: bool = False
    link_title: str = ""


class _PageReader(HTMLParser):
    """A Confluence page, its HTML export or its storage format, read into a `_Writer`.

    State rather than a tree: the open elements and what each changed, so a missing end tag closes
    what it should and a stray one closes nothing.
    """

    def __init__(self, *, export: bool) -> None:
        super().__init__(convert_charrefs=True)
        self.writer = _Writer()
        self.export = export
        self.reading = not export
        self.skip = 0
        self.open: list[_Open] = []
        self.buffer: list[str] = []
        self.level: int | None = None
        self.lists: list[str] = []
        self.pending: tuple[str, int] | None = None
        self.rows: list[list[str]] | None = None
        self.cell: list[str] | None = None
        self.tables = 0
        self.pre: list[str] | None = None
        self.capture: str = ""
        self.titles: dict[str, str] = {}

    # ---------------------------------------------------------------- the words
    def flush(self) -> None:
        words = " ".join("".join(self.buffer).split())
        self.buffer.clear()
        if not words:
            return
        if self.cell is not None:
            self.cell.append(words)
        elif self.level is not None:
            self.writer.heading(self.level, words)
        elif self.pending is not None:
            kind, level = self.pending
            self.pending = None
            if kind == "ol":
                self.writer.numbered(level, words)
            else:
                self.writer.bullet(level, words)
        elif self.lists:
            self.writer.continued(len(self.lists) - 1, words)
        else:
            self.writer.paragraph(words)

    def handle_data(self, data: str) -> None:
        if self.capture:
            self.titles[self.capture] = self.titles.get(self.capture, "") + data
        if self.skip or not self.reading:
            return
        if self.pre is not None:
            self.pre.append(data)
        else:
            self.buffer.append(data)

    def unknown_decl(self, data: str) -> None:
        # Python 3.13 hands a CDATA section here, which the storage format keeps code and link
        # text in.
        if data.startswith("CDATA["):
            self.handle_data(data[len("CDATA[") :])

    def handle_comment(self, data: str) -> None:
        # Later Pythons hand the same section here as a comment when it is not in foreign content.
        if data.startswith("[CDATA[") and data.endswith("]]"):
            self.handle_data(data[len("[CDATA[") : -2])

    # ---------------------------------------------------------------- opening
    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self._start(tag, dict(attrs), void=True)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self._start(tag, dict(attrs), void=tag in VOID_TAGS)

    def _start(self, tag: str, attrs: dict[str, str | None], *, void: bool) -> None:
        opened = _Open(tag=tag)
        classes = set((attrs.get("class") or "").split())
        ident = attrs.get("id") or ""
        if tag == "title" or ident == "title-text":
            self.capture = "title-text" if ident == "title-text" else "title"
            opened.effects.append("capture")
        if self.export and not self.reading and ident == "main-content":
            self.reading = True
            opened.effects.append("main")
        elif self.skip or tag in SKIPPED_TAGS or classes & SKIPPED_CLASSES:
            self.skip += 1
            opened.effects.append("skip")
        elif self.reading:
            self._effects(tag, attrs, classes, opened)
        if void:
            self._close(opened)
        else:
            self.open.append(opened)

    def _effects(
        self, tag: str, attrs: dict[str, str | None], classes: set[str], opened: _Open
    ) -> None:
        if tag in BLOCK_TAGS:
            self.flush()
        if len(tag) == 2 and tag[0] == "h" and tag[1] in "123456":
            self.level = int(tag[1]) + 1
            opened.effects.append("heading")
        elif tag in ("ol", "ul", "ac:task-list"):
            self.lists.append("ol" if tag == "ol" else "ul")
            if tag == "ol":
                self.writer.restart_level(len(self.lists) - 1)
            opened.effects.append("list")
        elif tag in ("li", "ac:task"):
            self.pending = (self.lists[-1] if self.lists else "ul", max(len(self.lists) - 1, 0))
            opened.effects.append("item")
        elif tag == "table":
            self.tables += 1
            if self.tables == 1:
                self.rows = []
            opened.effects.append("table")
        elif tag == "tr" and self.rows is not None and self.tables == 1:
            self.rows.append([])
        elif tag in ("td", "th") and self.rows is not None and self.tables == 1:
            self.cell = []
            opened.effects.append("cell")
        elif tag == "pre":
            self.pre = []
            opened.effects.append("pre")
        elif tag == "div" and _is_panel(classes):
            self._panel(_panel_label(classes), opened)
        elif tag == "ac:structured-macro":
            opened.macro = (attrs.get("ac:name") or "").strip().lower()
            opened.effects.append("macro")
            if opened.macro in PANEL_LABELS:
                self._panel(PANEL_LABELS[opened.macro], opened)
        elif tag in ("ac:rich-text-body", "ac:plain-text-body"):
            self._macro_has_body(tag)
            if tag == "ac:plain-text-body" and self._macro() in CODE_MACROS:
                self.pre = []
                opened.effects.append("pre")
        elif tag == "ac:plain-text-link-body":
            link = self._innermost("ac:link")
            if link is not None:
                link.had_body = True
        elif tag == "ri:page":
            link = self._innermost("ac:link")
            if link is not None:
                link.link_title = attrs.get("ri:content-title") or ""
        elif tag == "ac:link":
            opened.effects.append("link")
        elif tag == "time" and attrs.get("datetime"):
            self.buffer.append(f" {attrs['datetime']} ")

    def _panel(self, label: str, opened: _Open) -> None:
        self.flush()
        self.writer.quote_depth += 1
        opened.effects.append("quote")
        if label:
            self.writer.paragraph(f"**{label}**")

    def _innermost(self, tag: str) -> _Open | None:
        return next((one for one in reversed(self.open) if one.tag == tag), None)

    def _macro(self) -> str:
        found = self._innermost("ac:structured-macro")
        return "" if found is None else found.macro

    def _macro_has_body(self, tag: str) -> None:
        found = self._innermost("ac:structured-macro")
        if found is not None:
            found.had_body = True

    # ---------------------------------------------------------------- closing
    def handle_endtag(self, tag: str) -> None:
        if not any(one.tag == tag for one in self.open):
            return
        while self.open:
            closing = self.open.pop()
            self._close(closing)
            if closing.tag == tag:
                return

    def _close(self, closing: _Open) -> None:
        if closing.tag in BLOCK_TAGS and self.reading and not self.skip:
            self.flush()
        for effect in reversed(closing.effects):
            self._undo(effect, closing)

    def _undo(self, effect: str, closing: _Open) -> None:
        match effect:
            case "capture":
                self.capture = ""
            case "main":
                self.flush()
                self.reading = False
            case "skip":
                self.skip -= 1
            case "heading":
                self.flush()
                self.level = None
            case "list":
                self.flush()
                self.lists.pop()
                self.pending = None
            case "item":
                self.flush()
                self.pending = None
            case "cell":
                self.flush()
                if self.rows is not None and self.cell is not None:
                    if not self.rows:
                        self.rows.append([])
                    self.rows[-1].append(" ".join(self.cell))
                self.cell = None
            case "table":
                self.tables -= 1
                if self.tables == 0 and self.rows is not None:
                    rows, self.rows = self.rows, None
                    self.writer.table(rows)
            case "pre":
                code = "".join(self.pre or []).strip(chr(10))
                self.pre = None
                if code.strip():
                    self.writer.code(code)
            case "quote":
                self.flush()
                self.writer.quote_depth -= 1
            case "macro":
                if not closing.had_body and closing.macro not in NAVIGATION_MACROS:
                    self.flush()
                    self.writer.note(f"Confluence macro {closing.macro or 'with no name'}")
            case "link":
                if not closing.had_body and closing.link_title:
                    self.buffer.append(closing.link_title)
            case _:
                pass


def _is_panel(classes: set[str]) -> bool:
    """An information macro or a panel in the export; a code block's panel is its code."""
    return bool(classes & {"confluence-information-macro", "panel"}) and "code" not in classes


def _panel_label(classes: set[str]) -> str:
    """The word a panel of the export is labelled with, from its class."""
    prefix = "confluence-information-macro-"
    for one in sorted(classes):
        if one.startswith(prefix) and one[len(prefix) :] in PANEL_LABELS:
            return PANEL_LABELS[one[len(prefix) :]]
    return ""


def _page_text(page: str) -> str:
    export = MAIN_CONTENT_RE.search(page) is not None
    reader = _PageReader(export=export)
    reader.feed(page)
    reader.close()
    reader.flush()
    body = reader.writer.text()
    title = " ".join((reader.titles.get("title-text") or reader.titles.get("title") or "").split())
    if export:
        title = SPACE_PREFIX_RE.sub("", title, count=1)
    if not title:
        return body
    return f"# {title}{chr(10) * 2}{body}".strip()
