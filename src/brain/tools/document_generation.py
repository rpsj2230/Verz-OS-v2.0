"""The one tool that fills an approved company template and keeps the result as a Word file.

M12.3.5 asks that an agent produce a document from an approved company template through one
registered generation tool, and that the file be stored as an artifact holding only what the caller
may see. Every part of that but the tool already existed: a knowledge item has a kind (`template`)
and a verification, `brain.ops.artifact_store.StoredArtifacts.keep` records a produced file in the
name of the person it was produced for, and `python-docx` is already a dependency because the
parser reads Word with it. What nothing did was put the three behind a tool a run could call.

**What an approved template is.** A knowledge item of kind Template, published (not a draft, not
superseded, not archived), carrying a named person's verification, and **readable at the run's
reach**. A template nobody vouched for is never used, which is needs-rupash 193's own sentence.
A verification past its review date still counts: `brain.knowledge.item.badge` says "due for
review" and not "unverified", and refusing work on the day a review lapses would make the date a
switch somebody did not know they had.

**The file holds the template and the words the agent supplied, and nothing it did not read.**
There are two walls and each is enough alone. The template's passages come through
`brain.knowledge.document_tools.reader`, which applies the caller's reach inside every statement
and again at the table's row-level security, so a template the person cannot read is not read. The
values are text the agent wrote for named placeholders, and an agent's run is only ever shown what
its own reach returned, so there is no field in this tool by which a record could be named and
pulled in at a wider reach than the person asking. A tool that took a record id and a field would
have been a second reader of the company's data, with a second opinion about masks; rejected for
that reason, and the cost is that an agent copies a figure into a value after it has read it
through the tool that reads it.

**Nothing here computes a reach.** The run reach the gate hands every handler is already
`E(caller) ∩ agent_ceiling`, and it is both who the file is kept for and the reach the template is
read at. `brain.tools.skill_tools` makes the same remark for the same reason, and
`tests/invariants/test_single_implementation.py` holds the absence of a second implementation of
the rule.

**Refused as one sentence, and the sentence does not say which of five things was wrong.** A
template that does not exist, one the caller may not read, one that is not a template, one nobody
verified and one that is a draft are each `NO_APPROVED_TEMPLATE_IS_AVAILABLE`, because telling them
apart tells a person which documents exist (`brain.core.errors`). The one refusal that names
anything is the placeholders a readable, approved template still needs a value for, which are the
template's own words and which the person asking may already read.

**Classed as a read, which is a decision and is said here.** `SideEffect.NONE` is what makes a run
offer the tool at all (`brain.gate.runtime.AgentRuntime.offered` offers a model nothing that
changes something unless a connector prepares its action), and what it does is keep a file for the
person who asked, inside the product. Nothing is sent, nothing existing is changed and nothing
leaves the install, which is the website check's and the skill script's classification too.
The cost is that no approval stands between a run and a new file, and the file is the person's own
and is listed beside the others they were produced; a send of it is a different tool with its own
approval.

**Word only, and the format is a closed field.** The owner has not answered which formats follow
(needs-rupash 193): a slide deck and a spreadsheet each need a library the product does not hold,
and adding one is a decision about code that runs on every install. `DocumentFormat` has one
member, so a request for another is refused by the schema before anything runs, and the day an
answer arrives a second member is a reviewed change here.

**The layout is rebuilt from the template's text, and that is the limit.** The library keeps a
document's text as chunks and keeps the original file under its digest with no link back from the
item, so the original's styles, headers and logo are not reachable from here. Headings, paragraphs
and pipe tables are rebuilt from the text (a heading is a line the chunk's own section names, or a
Markdown heading), which reproduces what was written and not how it was set. Linking an item to its
original is a column on the item and a migration, and is the smallest thing that would carry a
company's own styling.

**The switch ships off and is asked at the call, not at start.** The registry is built once, so a
switch read when it was built would take a restart to follow; `generation_is_on` is asked by every
call and a tool that is switched off answers as an unavailable tool does. See
`brain.ops.features_declared.document_generation`.

**The artifact is attributed to the run and to the agent as it was.** The version recorded is a
digest of the agent's record without its lifecycle dates (`agent_version_of`), because the record
has no version number and a moving pin names nothing. The run is the request's trace.

Not done and said: an automation has no agent bound to its tool calls today (`RegistryToolCaller`
binds none), so a scheduled run is refused here rather than producing a file attributed to nobody.

Task ids: M12.3.5
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from collections.abc import Awaitable, Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from enum import StrEnum
from typing import Any, Final

import docx
import sqlalchemy as sa
import structlog
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from brain.agents.model import AgentRecord
from brain.console.agent_output import (
    Artifact,
    ArtifactError,
    ArtifactInput,
    ArtifactKind,
    Provenance,
)
from brain.core.entitlement import Capability, EntitlementSet, Grant
from brain.core.envelope import Entity, IdentityMode, SideEffect, ToolDefinition, TypedResult
from brain.core.errors import Denied
from brain.core.field_policy import Classification
from brain.knowledge.document_tools import (
    KNOWLEDGE_TOOL_PREFIX,
    MAX_PASSAGES,
    DocumentRead,
    KnowledgePassage,
    reach_through,
    reader,
)
from brain.knowledge.item import ITEM_ID_PATTERN, KnowledgeState, VerificationState, badge
from brain.knowledge.kinds import KnowledgeKind
from brain.knowledge.rows import RowQuery, RowSource
from brain.knowledge.search import KNOWLEDGE_READ, Reach, session_settings
from brain.ops.artifact_store import Produced, StoredArtifacts
from brain.ops.features import is_on
from brain.ops.features_declared.document_generation import DOCUMENT_GENERATION
from brain.ops.retention import DataClass
from brain.ops.storage import StorageError
from brain.tables.knowledge import KnowledgeItemRow
from brain.tools.registry import ResultContract, ToolRegistry

log = structlog.get_logger(__name__)

# ------------------------------------------------------------------ written-down reasons
#: Why every way a template can be unusable is one sentence.
NO_APPROVED_TEMPLATE_IS_AVAILABLE: Final = (
    "no approved template by that reference is available to this agent"
)

#: Why the tool answers like this when it is off, unbound or has nowhere to keep a file.
DOCUMENT_GENERATION_IS_NOT_AVAILABLE: Final = "document generation is not available to this run"

#: Why a template that is not approved is refused, and why the refusal does not say which way.
A_TEMPLATE_NOBODY_VOUCHED_FOR_IS_NEVER_FILLED: Final = (
    "Only a published knowledge item of kind Template that a named person has verified is filled, "
    "and only when the person asking may read it. Every other case, a draft, a withdrawn or "
    "replaced version, an unverified document, a different kind and a document they may not read, "
    "is answered with the one sentence, so that the answer does not say which documents exist."
)

#: Why a document holds only what its caller could see.
A_GENERATED_FILE_HOLDS_THE_TEMPLATE_AND_THE_AGENTS_WORDS_AND_NOTHING_ELSE: Final = (
    "The file is made of a template read at the run's reach and of text the agent wrote for named "
    "placeholders, and the agent was only ever shown what the same reach returned. There is no "
    "way in this tool to name a record and have it read at a wider reach, so the file cannot hold "
    "more than the person asking could read."
)

#: Why a value is inserted once and never read for placeholders of its own.
A_VALUE_IS_TEXT_AND_NEVER_A_TEMPLATE: Final = (
    "A value is inserted into the template in one pass. A value that itself contains a "
    "placeholder is text, and is written as it is: expanding it would let one value pull another "
    "into a place the template never named."
)

# ------------------------------------------------------------------ the vocabulary
#: The first segment of the tool's name, and the object its result is tagged with.
DOCUMENTS_SOURCE: Final = "documents"
GENERATED_DOCUMENT_OBJECT: Final = "generated_document"

#: What the tool is registered as.
GENERATE_DOCUMENT: Final = f"{DOCUMENTS_SOURCE}.generate_document"

#: What calling it needs. `invoke:` for the same reason the skill script's is: it is a permission to
#: act, and a field rule may not ask for one, so its result is read under `read:` (`own_results`).
GENERATE_CAPABILITY: Final = Capability(value=f"invoke:{GENERATED_DOCUMENT_OBJECT}")
GENERATED_DOCUMENT_READ: Final = Capability(value=f"read:{GENERATED_DOCUMENT_OBJECT}")

TOOL_DESCRIPTION: Final = (
    "Fill an approved company template with the words you supply for its placeholders and keep the "
    "result as a Word file for the person you are answering"
)

#: What the media type of the kept file is.
WORD_CONTENT_TYPE: Final = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

#: A placeholder as a template writes it: `{{name}}`, lower case, digits and underscores.
PLACEHOLDER: Final = re.compile(r"\{\{\s*([a-z][a-z0-9_]{0,63})\s*\}\}")

#: A name a value may be supplied under.
VALUE_NAME: Final = re.compile(r"^[a-z][a-z0-9_]{0,63}$")

#: What a request may carry. Bounded because a model chooses them.
MAX_VALUES: Final = 100
MAX_VALUE_CHARS: Final = 20_000
MAX_TOTAL_VALUE_CHARS: Final = 200_000
MAX_TITLE_CHARS: Final = 200

#: Characters an XML document cannot hold, which the library refuses and a model can produce.
_NOT_XML: Final = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

#: A Markdown heading in a template's text.
_HEADING: Final = re.compile(r"^(#{1,3})\s+(\S.*)$")


class DocumentFormat(StrEnum):
    """The formats a document is produced in. One, and the closure is the mechanism.

    A slide deck and a spreadsheet are waiting on needs-rupash 193, because each needs a library
    the product does not hold. A second member is a reviewed change to this enum and to
    `render`, and until then a request for another format fails the schema before it runs.
    """

    WORD = "word"


class GenerateDocument(BaseModel):
    """What a model may ask for: a template by its reference, a title and the placeholder values."""

    model_config = ConfigDict(frozen=True, extra="forbid")

    template_id: str = Field(pattern=ITEM_ID_PATTERN)
    title: str = Field(min_length=1, max_length=MAX_TITLE_CHARS)
    values: dict[str, str] = Field(default_factory=dict, max_length=MAX_VALUES)
    format: DocumentFormat = DocumentFormat.WORD

    @field_validator("values")
    @classmethod
    def _values_are_named_and_bounded(cls, v: dict[str, str]) -> dict[str, str]:
        """A name no placeholder could carry is refused, and so is a value too long to be a field.

        The total is bounded as well as each value, because a hundred values of the longest length
        is a request nobody wrote by hand.
        """
        for name, value in v.items():
            if not VALUE_NAME.match(name):
                msg = f"{name!r} is not a placeholder name"
                raise ValueError(msg)
            if len(value) > MAX_VALUE_CHARS:
                msg = f"the value for {name!r} is longer than a field of a document may be"
                raise ValueError(msg)
        if sum(len(one) for one in v.values()) > MAX_TOTAL_VALUE_CHARS:
            msg = "the values together are longer than a document's fields may be"
            raise ValueError(msg)
        return v


class GeneratedDocument(Entity):
    """What a run is told of the file it made: its reference, name, format and size.

    The file's content is not here and has no field to arrive in, so a document cannot come back
    through a result and past the artifact store's own reach check.
    """

    artifact_id: str
    title: str
    format: str
    size_bytes: int


@dataclass(frozen=True)
class RunIdentity:
    """Which agent, as it was, and which run a call belongs to. Bound by the caller of the tool."""

    agent_id: str
    agent_version: str
    run_id: str


def agent_version_of(record: AgentRecord) -> str:
    """A version for an agent record, which has none: a digest of what built the file.

    Its lifecycle dates are left out, because disabling and enabling an agent changes neither what
    it is nor what it made, and a version that moved with them would read as a different agent.
    """
    shown = record.model_dump(mode="json", exclude={"disabled_at", "archived_at"})
    digest = hashlib.sha256(json.dumps(shown, sort_keys=True).encode("utf-8")).hexdigest()
    return f"sha256:{digest[:32]}"


def identity_of(record: AgentRecord, run_id: str) -> RunIdentity:
    """The identity a run's calls of this tool are attributed to."""
    return RunIdentity(
        agent_id=record.agent_id, agent_version=agent_version_of(record), run_id=run_id
    )


# ------------------------------------------------------------------ is a template approved
@dataclass(frozen=True)
class TemplateFacts:
    """What the library records of one item that decides whether it may be filled.

    Satisfies `brain.knowledge.item.Vouched`, so the badge that is shown beside an answer is the
    reading used here, and the two cannot disagree about whether somebody vouched.
    """

    kind: str | None
    state: KnowledgeState
    verified_by: str
    verified_at: datetime | None
    review_by: datetime | None


FACTS_COLUMNS: Final[tuple[str, ...]] = ("kind", "state", "verified_by", "verified_at", "review_by")


def facts_query(item_id: str, *, reach: Reach) -> RowQuery:
    """One item's kind, state and verification, under the caller's session settings.

    Read from `know.item`, whose row-level security is the second wall, so an item the person
    cannot read returns no row at all.
    """
    statement = (
        sa.select(*(getattr(KnowledgeItemRow, name).label(name) for name in FACTS_COLUMNS))
        .where(KnowledgeItemRow.item_id == item_id)
        .limit(1)
    )
    return RowQuery(
        entity="knowledge_item",
        source=KNOWLEDGE_TOOL_PREFIX,
        columns=FACTS_COLUMNS,
        statement=statement,
        certainly_empty=False,
        settings=session_settings(reach),
    )


def facts_of(row: Mapping[str, Any]) -> TemplateFacts | None:
    """A row as facts, or None when it does not say what an item's state is."""
    try:
        state = KnowledgeState(str(row["state"]))
    except ValueError:
        return None
    return TemplateFacts(
        kind=None if row["kind"] is None else str(row["kind"]),
        state=state,
        verified_by="" if row["verified_by"] is None else str(row["verified_by"]),
        verified_at=row["verified_at"],
        review_by=row["review_by"],
    )


def is_approved_template(facts: TemplateFacts, *, now: datetime) -> bool:
    """Whether this item is a published template somebody verified. See the module docstring.

    Verified, or verified and past its review: the badge's own two readings of vouched-for.
    """
    if facts.kind != KnowledgeKind.TEMPLATE.value or facts.state is not KnowledgeState.PUBLISHED:
        return False
    return badge(facts, now=now).state in (VerificationState.VERIFIED, VerificationState.DUE)


# ------------------------------------------------------------------ the layout
class BlockType(StrEnum):
    HEADING = "heading"
    PARAGRAPH = "paragraph"
    TABLE = "table"


@dataclass(frozen=True)
class Block:
    """One heading, paragraph or table of a template, before any value is put in it.

    `level` is a heading's depth. `lines` is a paragraph's lines, or a heading's one line, and
    `rows` is a table's cells.
    """

    kind: BlockType
    lines: tuple[str, ...] = ()
    rows: tuple[tuple[str, ...], ...] = ()
    level: int = 1


def _cells(line: str) -> tuple[str, ...]:
    """A pipe row's cells, as `brain.knowledge.text_path.pipe_row` wrote them."""
    inner = line.strip().removeprefix("|").removesuffix("|")
    return tuple(one.strip() for one in inner.split(" | "))


class _Layout:
    """The blocks being built, and the paragraph and table still open."""

    def __init__(self) -> None:
        self.blocks: list[Block] = []
        self._paragraph: list[str] = []
        self._table: list[tuple[str, ...]] = []

    def close(self) -> None:
        """End whatever is open: a blank line, a heading and the end of a passage all do."""
        if self._paragraph:
            self.blocks.append(Block(kind=BlockType.PARAGRAPH, lines=tuple(self._paragraph)))
            self._paragraph = []
        if self._table:
            self.blocks.append(Block(kind=BlockType.TABLE, rows=tuple(self._table)))
            self._table = []

    def heading(self, text: str, level: int) -> None:
        self.close()
        self.blocks.append(Block(kind=BlockType.HEADING, lines=(text,), level=level))

    def row(self, line: str) -> None:
        if self._paragraph:
            self.close()
        self._table.append(_cells(line))

    def text(self, line: str) -> None:
        if self._table:
            self.close()
        self._paragraph.append(line)


def layout(passages: Sequence[KnowledgePassage]) -> tuple[Block, ...]:
    """A template's passages, in reading order, as blocks.

    A line the passage's own section names, when it opens the passage, is a heading, which is how a
    Word heading survives being parsed to text; a Markdown heading is one too; a run of lines that
    each begin with a pipe is a table; everything else is a paragraph, and a blank line ends one.
    """
    built = _Layout()
    for passage in passages:
        opening = True
        for raw in passage.document.split("\n"):
            line = raw.rstrip()
            if not line.strip():
                built.close()
                continue
            marked = _HEADING.match(line.strip())
            if line.lstrip().startswith("|"):
                built.row(line)
            elif marked is not None:
                built.heading(marked.group(2).strip(), len(marked.group(1)))
            elif opening and passage.section and line.strip() == passage.section.strip():
                built.heading(line.strip(), 1)
            else:
                built.text(line)
            opening = False
        built.close()
    return tuple(built.blocks)


def placeholders_in(blocks: Sequence[Block]) -> tuple[str, ...]:
    """Every placeholder a template names, once each, in the order it first names one."""
    found: dict[str, None] = {}
    for block in blocks:
        texts = [*block.lines, *(cell for row in block.rows for cell in row)]
        for text in texts:
            for name in PLACEHOLDER.findall(text):
                found.setdefault(name)
    return tuple(found)


def filled(text: str, values: Mapping[str, str]) -> str:
    """One piece of text with its placeholders replaced, in one pass.

    See `A_VALUE_IS_TEXT_AND_NEVER_A_TEMPLATE`. A placeholder with no value is left as written,
    which `missing_values` has already refused before a file is made.
    """
    return PLACEHOLDER.sub(lambda found: values.get(found.group(1), found.group(0)), text)


def _clean(text: str) -> str:
    """Text a Word file can hold: the characters an XML document cannot are dropped."""
    return _NOT_XML.sub("", text)


def render(
    title: str, blocks: Sequence[Block], values: Mapping[str, str], *, now: datetime
) -> bytes:
    """The Word file: the title, then every block with its values put in.

    The author and last-modified-by are blank and the dates are the run's own, because the
    library's template carries a name and a date of its own that would read as a person.
    """
    document = docx.Document()
    properties = document.core_properties
    properties.title = _clean(title)[:MAX_TITLE_CHARS]
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = now
    properties.modified = now
    document.add_heading(_clean(title), level=0)
    for block in blocks:
        match block.kind:
            case BlockType.HEADING:
                document.add_heading(_clean(filled(block.lines[0], values)), level=block.level)
            case BlockType.PARAGRAPH:
                document.add_paragraph(_clean(filled("\n".join(block.lines), values)))
            case BlockType.TABLE:
                width = max(len(row) for row in block.rows)
                table = document.add_table(rows=len(block.rows), cols=width)
                table.style = "Table Grid"
                for index, row in enumerate(block.rows):
                    for column, cell in enumerate(row):
                        table.cell(index, column).text = _clean(filled(cell, values))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def missing_values(blocks: Sequence[Block], values: Mapping[str, str]) -> tuple[str, ...]:
    """The placeholders a template names that the request gave no value for, in order."""
    return tuple(name for name in placeholders_in(blocks) if name not in values)


# ------------------------------------------------------------------ the tool
#: Where a produced file is kept, as a callable and not a protocol: a protocol method here would be
#: a door to the outside that `brain.ops.effects` asks the classification of, and keeping a file
#: twice is two files, which is the answer `StoredArtifacts.keep` already gives and records.
Keep = Callable[[Produced, datetime], Awaitable[Artifact]]


def keeping(store: StoredArtifacts) -> Keep:
    """`StoredArtifacts.keep`, as the callable the tool is bound to."""

    async def keep(produced: Produced, at: datetime) -> Artifact:
        return await store.keep(produced, at=at)

    return keep


@dataclass(frozen=True)
class DocumentTools:
    """What the tool is bound to: where templates are read, where files are kept, and the switch.

    `is_on` is asked by every call. All three are required and none has a default, for the reason
    `brain.tools.run_skill.SkillScriptTool` gives about its runner: a tool bound to nothing would
    be present and unable to answer.
    """

    records: RowSource
    keep: Keep
    is_on: Callable[[], Awaitable[bool]]


async def generation_is_on(sessions: async_sessionmaker[AsyncSession]) -> bool:
    """Whether this install has switched document generation on. Off until somebody does."""
    async with sessions() as session:
        return await is_on(session, DOCUMENT_GENERATION)


def definition() -> ToolDefinition:
    """What the catalogue describes to a model.

    `IdentityMode.DELEGATED` because the tool spends no credential of its own: it reads at the
    run's reach and keeps in the caller's name. `SideEffect.NONE` is argued in the module docstring.
    """
    return ToolDefinition(
        name=GENERATE_DOCUMENT,
        description=TOOL_DESCRIPTION,
        entity=GENERATED_DOCUMENT_OBJECT,
        args_schema=GenerateDocument.model_json_schema(),
        required_capability=GENERATE_CAPABILITY.value,
        side_effect=SideEffect.NONE,
        identity_mode=IdentityMode.DELEGATED,
        source=DOCUMENTS_SOURCE,
    )


def inputs_of(template_id: str) -> tuple[ArtifactInput, ...]:
    """What the file was made from, for the class it is kept under.

    The template is knowledge, and the agent's words are a run's payload, which is the short one:
    `brain.console.agent_output.retention_class_for` takes the narrowest window, so the file is
    labelled with the shortest-lived thing it was built from. `template_id` labels the input and
    is not shown to a reader.
    """
    return (
        ArtifactInput(
            label=f"template {template_id}",
            data_class=DataClass.LEARNED,
            classification=Classification.INTERNAL,
        ),
        ArtifactInput(
            label="words the agent supplied",
            data_class=DataClass.PAYLOAD,
            classification=Classification.INTERNAL,
        ),
    )


def handler(
    tools: DocumentTools,
) -> Callable[..., Awaitable[TypedResult[GeneratedDocument]]]:
    """The callable a registry registers, called as every dispatcher here calls a handler."""
    read = reader(tools.records)

    async def generate_document(
        request: GenerateDocument,
        *,
        entitlement: EntitlementSet,
        now: datetime | None = None,
        produced_for: RunIdentity | None = None,
    ) -> TypedResult[GeneratedDocument]:
        """Read the template at the run's reach, fill it, keep the file, and say which it is."""
        moment = now if now is not None else datetime.now(UTC)
        if produced_for is None or not await tools.is_on():
            raise Denied(DOCUMENT_GENERATION_IS_NOT_AVAILABLE)
        reach = await reach_through(tools.records, entitlement, moment)
        if reach is None:
            raise Denied(NO_APPROVED_TEMPLATE_IS_AVAILABLE)
        found = await tools.records.rows(facts_query(request.template_id, reach=reach))
        facts = facts_of(found[0]) if found else None
        if facts is None or not is_approved_template(facts, now=moment):
            raise Denied(NO_APPROVED_TEMPLATE_IS_AVAILABLE)
        template = await read(
            DocumentRead(document_id=request.template_id, limit=MAX_PASSAGES),
            entitlement=entitlement,
            now=moment,
        )
        # A template as long as the read's limit may have been cut, and a file made from half a
        # template is a wrong document with nothing to say so.
        if not template.records or template.truncated:
            raise Denied(NO_APPROVED_TEMPLATE_IS_AVAILABLE)
        blocks = layout(template.records)
        absent = missing_values(blocks, request.values)
        if absent:
            raise Denied("the template needs a value for: " + ", ".join(absent))
        body = render(request.title, blocks, request.values, now=moment)
        held = entitlement.scope_for(KNOWLEDGE_READ, moment)
        try:
            kept = await tools.keep(
                Produced(
                    agent_id=produced_for.agent_id,
                    kind=ArtifactKind.DOCUMENT,
                    run_id=produced_for.run_id,
                    agent_version=produced_for.agent_version,
                    caller_id=entitlement.principal_id,
                    reach=entitlement,
                    inputs=inputs_of(request.template_id),
                    body=body,
                    content_type=WORD_CONTENT_TYPE,
                    provenance=Provenance(knowledge_items=(request.template_id,)),
                    drew_on=() if held is None else (Grant(capability=KNOWLEDGE_READ, scope=held),),
                ),
                moment,
            )
        except (ArtifactError, StorageError) as failure:
            log.warning("document.not_kept", error=type(failure).__name__)
            raise Denied(DOCUMENT_GENERATION_IS_NOT_AVAILABLE) from failure
        return TypedResult[GeneratedDocument](
            records=(
                GeneratedDocument(
                    entity=GENERATED_DOCUMENT_OBJECT,
                    id=kept.artifact_id,
                    artifact_id=kept.artifact_id,
                    title=request.title,
                    format=request.format.value,
                    size_bytes=kept.bytes_stored,
                ),
            ),
            source=DOCUMENTS_SOURCE,
            fetched_at=moment.isoformat(),
        )

    return generate_document


def register_document_generation(registry: ToolRegistry, tools: DocumentTools) -> None:
    """Register the tool through the one door every tool uses, with its contract declared."""
    registry.register(definition(), handler(tools), result_contract=ResultContract.TYPED)
