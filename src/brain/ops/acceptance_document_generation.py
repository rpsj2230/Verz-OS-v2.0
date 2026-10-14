"""The install acceptance check for an agent filling an approved template into a Word file.

One check, over the product's own pieces inside the check's transaction (M12.3.5). A steward of
acceptance_a adds a document of kind Template and a member of acceptance_a is given an agent that
may use the one document tool; a member of acceptance_b holds the same capability and no more. The
tool is the one `brain.tools.startup.build_registry` registers, over the install's own rows and a
file store held in this process, called as the gate calls a handler: with the run reach, which is
the caller's narrowed by the agent's ceiling through the one intersection.

**What is checked.**
- The switch is pressed through the Features route both ways and never assumed: off, the tool is
  refused as an unavailable one and nothing is kept; on, it works.
- A template nobody has verified is refused with the one sentence, and so is a reference to a
  document the asker's colleague in acceptance_b holds, so the refusal does not say which exists.
- Once the steward verifies it, the file is made: it holds the template's words with the values put
  in, it is kept in the asker's name, the asker may fetch it and the other department's member may
  not, and a placeholder left without a value is refused by name with nothing kept.
- An agent whose ceiling does not name the tool is offered nothing by the catalogue, and one whose
  ceiling does is offered it.

**The file is kept in memory.** The artifact is written through the product's own store over an
object store held in this process, so nothing reaches a bucket and the rows are rolled back with the
check. What is proved is the tool, the template rules, the attribution and the download verdict, not
the bucket; the install's own object store is the artifact checks' and stays theirs.

Task ids: M12.3.5
"""

from __future__ import annotations

import io
from datetime import timedelta
from functools import partial
from typing import TYPE_CHECKING, Any, Final, cast

from brain.core.scope import Scope
from brain.ops.acceptance import RESERVED_DEPARTMENTS, CheckFailedError, check
from brain.ops.acceptance_operations_console import asking_as, console_for, traced
from brain.ops.acceptance_run import Harness

if TYPE_CHECKING:
    from collections.abc import Awaitable, Callable

    from brain.core.envelope import TypedResult
    from brain.tools.document_generation import GeneratedDocument

#: Where this module's checks stand on the Install page, before every larger key. See
#: `brain.ops.acceptance.A_CHECK_MODULE_IS_FOUND_AND_PLACES_ITSELF`.
CHECK_ORDER: Final = 2350

A, B = RESERVED_DEPARTMENTS

#: The authority that presses a feature's switch. Restated rather than imported, so a change to the
#: route's authority fails this check instead of moving with it.
SWITCHES_FEATURES: Final = "admin:feature"

#: A template's words. The placeholders are two, and the check gives each a value of its own.
TEMPLATE_BODY: Final = (
    b"# Acceptance proposal\n\nPrepared for {{client_name}} by {{author_name}}.\n"
)

OFF_AND_MADE: Final = "a file was produced while document generation was switched off"
UNVERIFIED_WAS_FILLED: Final = "a template nobody had verified was filled"
OTHERS_WAS_FILLED: Final = "a document the asker's own reach does not cover was filled"
NOT_MADE: Final = "an approved template was not filled for the person who asked"
WRONG_WORDS: Final = "the file did not hold the template's words with the values put in"
NOT_THEIRS: Final = "the file was not kept in the name of the person who asked"
OTHERS_MAY_FETCH: Final = "a member of another department was let fetch the file"
ASKER_MAY_NOT_FETCH: Final = "the person the file was made for could not fetch it"
HOLE_FILLED: Final = "a placeholder left without a value was filled or the refusal named a value"
HOLE_KEPT: Final = "a file was kept although a placeholder had no value"
CATALOGUE_WIDE: Final = "an agent whose ceiling does not name the tool was offered it"
CATALOGUE_NARROW: Final = "an agent whose ceiling names the tool was not offered it"
FILE_NOT_KEPT: Final = "an approved template produced a file nothing recorded"


class _Shelf:
    """An object store held in memory: a produced file is kept nowhere but this process."""

    def __init__(self) -> None:
        self.kept: dict[str, bytes] = {}

    def put_object(self, bucket_name: str, key: str, body: bytes, content_type: str) -> None:
        self.kept[f"{bucket_name}/{key}"] = body

    def get_object(self, bucket_name: str, key: str) -> bytes:
        return self.kept[f"{bucket_name}/{key}"]

    def delete_object(self, bucket_name: str, key: str) -> None:
        self.kept.pop(f"{bucket_name}/{key}", None)

    def list_objects(self, bucket_name: str, prefix: str) -> Any:
        return iter(())


async def _pressed(h: Harness, switcher: str, *, on: bool, n: int) -> None:
    """The feature's switch, through the Features route, as the administrator presses it."""
    from brain.feature_routes import SwitchAsked, switch_feature
    from brain.ops.features_declared.document_generation import DOCUMENT_GENERATION

    async with traced(h, n):
        await switch_feature(
            console_for(h).request("POST"),
            DOCUMENT_GENERATION.name,
            SwitchAsked(on=on),
            await asking_as(h, switcher, strong=True),
        )


@check(
    leaves=("M12.3.5",),
    sentence=(
        "An agent's run fills a template its steward has verified into a Word file kept in the "
        "asker's name, which they may fetch and another department's member may not; an unverified "
        "or another department's template is refused alike, a missing value keeps nothing, and "
        "the tool is refused while the switch is off or to an agent whose ceiling omits it."
    ),
)
async def an_approved_template_is_filled_for_its_asker_alone(
    h: Harness,
) -> None:
    import docx

    from brain.agents.model import tool_ceiling
    from brain.console.agent_output import may_download
    from brain.console.workspace_capabilities import run_reach
    from brain.core.errors import Denied
    from brain.gate.catalogue import project
    from brain.gate.roster import run_entitlement
    from brain.knowledge.kinds import KnowledgeKind
    from brain.knowledge.row_store import SessionRowSource
    from brain.ops.acceptance_checks import KNOWLEDGE_READS, _in, _upload
    from brain.ops.acceptance_checks_lifecycle import _as, _verify
    from brain.ops.acceptance_workspace import installed_agent, stored_agent
    from brain.ops.artifact_store import StoredArtifacts
    from brain.tools.document_generation import (
        DOCUMENT_GENERATION_IS_NOT_AVAILABLE,
        GENERATE_CAPABILITY,
        GENERATE_DOCUMENT,
        GENERATED_DOCUMENT_READ,
        NO_APPROVED_TEMPLATE_IS_AVAILABLE,
        DocumentTools,
        GenerateDocument,
        RunIdentity,
        agent_version_of,
        generation_is_on,
        keeping,
    )
    from brain.tools.startup import build_registry

    await h.found_departments()
    library, library_b, asker, outsider, switcher = (
        h.principal(A, "library"),
        h.principal(B, "library"),
        h.principal(A, "member"),
        h.principal(B, "member"),
        h.principal(A, "switcher"),
    )
    await h.person(library, department=A, grants=_in(A, "admin:knowledge"))
    await h.person(library_b, department=B, grants=_in(B, "admin:knowledge"))
    await h.person(asker, department=A, grants=_in(A, *KNOWLEDGE_READS, GENERATE_CAPABILITY.value))
    await h.person(
        outsider, department=B, grants=_in(B, *KNOWLEDGE_READS, GENERATE_CAPABILITY.value)
    )
    await h.person(switcher, department=A, grants=((SWITCHES_FEATURES, Scope.unrestricted()),))

    async def template_of(owner: str, body: bytes) -> str:
        made = await _upload(
            h,
            owner,
            filename=f"Acceptance template {h.word()}.md",
            declared="text/markdown",
            body=body,
            kind=KnowledgeKind.TEMPLATE,
            department=B if owner == library_b else A,
        )
        if not hasattr(made, "item"):
            raise CheckFailedError("a well-formed Markdown template could not be added")
        return str(made.item.item_id)

    ours = await template_of(library, TEMPLATE_BODY)
    theirs = await template_of(library_b, TEMPLATE_BODY)

    good_agent = await installed_agent(
        h,
        library,
        capabilities=(*KNOWLEDGE_READS, GENERATE_CAPABILITY.value, GENERATED_DOCUMENT_READ.value),
        allowed_tools=(GENERATE_DOCUMENT,),
        suffix="_documents",
    )
    bare_agent = await installed_agent(
        h, library, capabilities=KNOWLEDGE_READS, allowed_tools=(), suffix="_documents_bare"
    )
    good, bare = await stored_agent(h, good_agent), await stored_agent(h, bare_agent)

    shelf = _Shelf()
    store = StoredArtifacts(h.sessions, shelf, "acceptance")
    records = SessionRowSource(h.sessions)
    registry = build_registry(
        source=h.settings.tool_source,
        records=records,
        documents=DocumentTools(
            records=records, keep=keeping(store), is_on=partial(generation_is_on, h.sessions)
        ),
    )
    # The registry keeps a handler as an untyped callable; this is the one it registered.
    call = cast(
        "Callable[..., Awaitable[TypedResult[GeneratedDocument]]]",
        registry.get(GENERATE_DOCUMENT).handler,
    )
    caller = await h.reach(asker)
    run_reach_of = run_reach(caller, good)
    identity = RunIdentity(
        agent_id=good_agent, agent_version=agent_version_of(good), run_id=h.trace_id
    )
    values = {"client_name": f"{h.word()} Ltd", "author_name": h.word()}

    async def asking(template: str, given: dict[str, str]) -> Any:
        return await call(
            GenerateDocument(template_id=template, title="Acceptance proposal", values=given),
            entitlement=run_reach_of,
            now=h.now,
            produced_for=identity,
        )

    async def sentence_of(template: str, given: dict[str, str]) -> str | None:
        """The sentence a call was refused with, or None if it was not refused."""
        try:
            await asking(template, given)
        except Denied as refused:
            return str(refused)
        return None

    # The agent's catalogue: the tool is there for the one whose ceiling names it, and only there.
    shown = {
        name: project(
            registry, run_entitlement(caller, record), tool_ceiling(record), now=h.now
        ).names
        for name, record in (("good", good), ("bare", bare))
    }
    if GENERATE_DOCUMENT not in shown["good"]:
        raise CheckFailedError(CATALOGUE_NARROW)
    if GENERATE_DOCUMENT in shown["bare"]:
        raise CheckFailedError(CATALOGUE_WIDE)

    # The switch, both ways, whatever state the install was in.
    await _pressed(h, switcher, on=False, n=1)
    if await sentence_of(ours, values) != DOCUMENT_GENERATION_IS_NOT_AVAILABLE:
        raise CheckFailedError(OFF_AND_MADE)
    await _pressed(h, switcher, on=True, n=2)

    # A template nobody verified, and one that is another department's: the one sentence.
    if await sentence_of(ours, values) != NO_APPROVED_TEMPLATE_IS_AVAILABLE:
        raise CheckFailedError(UNVERIFIED_WAS_FILLED)
    steward = await _as(h, library)
    if await _verify(h, steward, ours, review_by=h.now + timedelta(days=30)) is None:
        raise CheckFailedError("the steward of an uploaded template could not verify it")
    steward_b = await _as(h, library_b)
    if await _verify(h, steward_b, theirs, review_by=h.now + timedelta(days=30)) is None:
        raise CheckFailedError("the steward of an uploaded template could not verify it")
    if await sentence_of(theirs, values) != NO_APPROVED_TEMPLATE_IS_AVAILABLE:
        raise CheckFailedError(OTHERS_WAS_FILLED)

    # A placeholder with no value is refused, and nothing is kept for it.
    held_back = await sentence_of(ours, {"client_name": values["client_name"]})
    if held_back is None or "author_name" not in held_back or "client_name" in held_back:
        raise CheckFailedError(HOLE_FILLED)
    if shelf.kept:
        raise CheckFailedError(HOLE_KEPT)

    # The positive case: filled, kept in the asker's name, fetched by them and by nobody in B.
    made = await asking(ours, values)
    if len(made.records) != 1:
        raise CheckFailedError(NOT_MADE)
    kept = await store.one(made.records[0].artifact_id)
    if kept is None or len(shelf.kept) != 1:
        raise CheckFailedError(FILE_NOT_KEPT)
    body = await store.bytes_of(kept)
    document = docx.Document(io.BytesIO(body))
    text = "\n".join(one.text for one in document.paragraphs)
    if values["client_name"] not in text or values["author_name"] not in text or "{{" in text:
        raise CheckFailedError(WRONG_WORDS)
    if kept.artifact.caller_id != asker or kept.artifact.agent_id != good_agent:
        raise CheckFailedError(NOT_THEIRS)
    if not may_download(kept.artifact, caller, h.now):
        raise CheckFailedError(ASKER_MAY_NOT_FETCH)
    if may_download(kept.artifact, await h.reach(outsider), h.now):
        raise CheckFailedError(OTHERS_MAY_FETCH)
